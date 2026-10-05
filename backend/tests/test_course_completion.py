"""Course completion report: admin opens a period per semester, lecturers
submit covered/planned chapters until the deadline, admin reopens for chosen
lecturers, filters by semester/class and downloads Excel."""
import io
from datetime import datetime, time, timedelta

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models import CompletionPeriod, Course, Notification, Timetable
from app.utils.settings import get_setting_str
from tests.conftest import auth_header, make_user

FMT = "%Y-%m-%dT%H:%M"


@pytest.fixture()
def teaching(app, lecturer):
    db.session.add(Timetable(lecturer_id=lecturer.id, course_name="Structural Analysis", batch="BCE08",
                             days="Mon", start_time=time(8), end_time=time(10)))
    db.session.add(Course(name="Hydraulics", batch="BCE09", lecturer_id=lecturer.id))
    db.session.commit()


@pytest.fixture()
def other_lecturer(app):
    user = make_user("other@example.com", role="lecturer")
    db.session.add(Timetable(lecturer_id=user.id, course_name="Surveying", batch="BCE08",
                             days="Tue", start_time=time(8), end_time=time(10)))
    db.session.commit()
    return user


def open_period(client, admin, **dates):
    now = datetime.now()
    body = {"opens_at": (now - timedelta(hours=1)).strftime(FMT), "closes_at": (now + timedelta(days=4)).strftime(FMT)}
    body.update(dates)
    return client.post("/api/course-completion/periods", headers=auth_header(admin), json=body)


def save(client, lecturer, period_id, **overrides):
    body = {"period_id": period_id, "class_name": "BCE08", "course_name": "Structural Analysis",
            "covered_chapters": 6, "planned_chapters": 8}
    body.update(overrides)
    return client.put("/api/course-completion/mine", headers=auth_header(lecturer), json=body)


def test_closed_until_admin_opens(client, lecturer, teaching):
    view = client.get("/api/course-completion/mine", headers=auth_header(lecturer)).get_json()
    assert view["period"] is None and not view["can_submit"]
    assert save(client, lecturer, 1).status_code == 404


def test_lecturer_submits_and_edits(client, admin, lecturer, teaching):
    res = open_period(client, admin)
    assert res.status_code == 201, res.get_json()
    period = res.get_json()
    assert period["semester"] == get_setting_str("semester_name")
    # Lecturers are told it's open.
    assert Notification.query.filter_by(user_id=lecturer.id, type="course_completion").count() == 1

    view = client.get("/api/course-completion/mine", headers=auth_header(lecturer)).get_json()
    assert view["can_submit"]
    assert [(r["class_name"], r["course_name"]) for r in view["rows"]] == [("BCE08", "Structural Analysis"), ("BCE09", "Hydraulics")]

    res = save(client, lecturer, period["id"])
    assert res.status_code == 200, res.get_json()
    row = res.get_json()["rows"][0]
    assert row["completion_percent"] == 75.0 and row["lecturer_name"] == lecturer.name

    save(client, lecturer, period["id"], covered_chapters=8)
    view = client.get("/api/course-completion/mine", headers=auth_header(lecturer)).get_json()
    assert view["rows"][0]["completion_percent"] == 100.0
    assert CompletionPeriod.query.one().reports[0].covered_chapters == 8


@pytest.mark.parametrize("overrides", [
    {"covered_chapters": 9}, {"planned_chapters": 0}, {"covered_chapters": -1}, {"covered_chapters": ""},
    {"planned_chapters": "abc"},
])
def test_validation(client, admin, lecturer, teaching, overrides):
    pid = open_period(client, admin).get_json()["id"]
    assert save(client, lecturer, pid, **overrides).status_code == 400


def test_only_assigned_courses(client, admin, lecturer, teaching, other_lecturer):
    pid = open_period(client, admin).get_json()["id"]
    assert save(client, lecturer, pid, course_name="Surveying").status_code == 403


def test_deadline_closes_and_reopen_for_one_lecturer(client, admin, lecturer, teaching, other_lecturer):
    pid = open_period(client, admin).get_json()["id"]
    period = CompletionPeriod.query.get(pid)
    period.opens_at = datetime.now() - timedelta(days=5)
    period.closes_at = datetime.now() - timedelta(minutes=1)
    db.session.commit()

    assert save(client, lecturer, pid).status_code == 403
    view = client.get("/api/course-completion/mine", headers=auth_header(lecturer)).get_json()
    assert not view["can_submit"] and "closed" in view["message"]

    until = (datetime.now() + timedelta(days=2)).strftime(FMT)
    res = client.post(f"/api/course-completion/periods/{pid}/reopen", headers=auth_header(admin),
                      json={"lecturer_ids": [lecturer.id], "until": until})
    assert res.status_code == 200, res.get_json()
    assert res.get_json()["reopenings"][0]["active"]
    assert save(client, lecturer, pid).status_code == 200
    # Still closed for everyone else.
    assert save(client, other_lecturer, pid, course_name="Surveying").status_code == 403

    client.delete(f"/api/course-completion/periods/{pid}/reopen/{lecturer.id}", headers=auth_header(admin))
    assert save(client, lecturer, pid).status_code == 403


def test_close_now_and_period_rules(client, admin, lecturer, teaching):
    assert open_period(client, admin, closes_at=(datetime.now() - timedelta(hours=2)).strftime(FMT)).status_code == 400
    pid = open_period(client, admin).get_json()["id"]
    assert open_period(client, admin).status_code == 409  # one period per semester
    assert client.post(f"/api/course-completion/periods/{pid}/close", headers=auth_header(admin)).get_json()["status"] == "closed"
    assert save(client, lecturer, pid).status_code == 403
    # Next semester gets its own period.
    assert open_period(client, admin, semester="March 2027 - July 2027").status_code == 201


def test_permissions(client, admin, lecturer, teaching):
    assert open_period(client, lecturer).status_code == 403
    assert client.get("/api/course-completion/reports", headers=auth_header(lecturer)).status_code == 403
    assert client.get("/api/course-completion/periods", headers=auth_header(admin)).status_code == 200


def test_admin_reports_filter_and_excel(client, admin, lecturer, teaching, other_lecturer):
    pid = open_period(client, admin).get_json()["id"]
    save(client, lecturer, pid)
    save(client, lecturer, pid, class_name="BCE09", course_name="Hydraulics", covered_chapters=1, planned_chapters=3)

    data = client.get("/api/course-completion/reports", headers=auth_header(admin)).get_json()
    assert data["classes"] == ["BCE08", "BCE09"]
    assert len(data["rows"]) == 3  # Surveying is still missing
    by_class = client.get("/api/course-completion/reports?class_name=BCE08", headers=auth_header(admin)).get_json()
    assert {(r["course_name"], r["submitted"]) for r in by_class["rows"]} == {("Structural Analysis", True), ("Surveying", False)}

    res = client.get("/api/course-completion/export.xlsx", headers=auth_header(admin))
    ws = load_workbook(io.BytesIO(res.data)).active
    assert [c.value for c in ws[1]] == ["Semester", "Class Name", "Lecturer Name", "Course Name",
                                        "Covered Chapters", "Planned Chapters", "Completion %"]
    assert ws.max_row == 3
    assert ws.cell(row=2, column=3).value == lecturer.name and ws.cell(row=2, column=7).value == 0.75

    one = client.get("/api/course-completion/export.xlsx?class_name=BCE09", headers=auth_header(admin))
    ws = load_workbook(io.BytesIO(one.data)).active
    assert ws.max_row == 2 and ws.cell(row=2, column=4).value == "Hydraulics"
