"""Special exam registration: student submits (only inside the registration
period, optionally with evidence documents), Administration Team decides,
student is notified and emailed, staff export Excel."""
import io
from datetime import date, datetime, time, timedelta

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models import EmailOutbox, FacultyRole, Notification, SpecialExamPeriod, SpecialExamRequest, Timetable
from tests.conftest import auth_header, make_user

PDF = b"%PDF-1.4\n%test\n"


@pytest.fixture(autouse=True)
def registration_open(app, request):
    """Registration is closed by default; most tests run inside an open period."""
    if "closed_registration" in request.keywords:
        return None
    now = datetime.now()
    period = SpecialExamPeriod(opens_at=now - timedelta(days=1), closes_at=now + timedelta(days=2))
    db.session.add(period)
    db.session.commit()
    return period


@pytest.fixture()
def student(app):
    return make_user("student@example.com", role="student", batch="BCE08", student_id_number="S123")


@pytest.fixture()
def admin_team(app):
    user = make_user("team@example.com", role="lecturer")
    db.session.add(FacultyRole(user_id=user.id, role="admin_team"))
    db.session.commit()
    return user


@pytest.fixture()
def timetable(app, lecturer):
    db.session.add(Timetable(lecturer_id=lecturer.id, course_name="Structural Analysis", batch="BCE08",
                             days="Mon", start_time=time(8), end_time=time(10)))
    db.session.commit()


def valid(**overrides):
    body = {
        "full_name": "Amina Ali", "batch": "BCE08", "course_name": "Structural Analysis",
        "reason": "I was in hospital.", "exam_date": (date.today() - timedelta(days=3)).isoformat(),
        "shift": "1", "phone": "+252 61 5123456", "id_number": "S123",
    }
    body.update(overrides)
    return body


def submit(client, student, **overrides):
    return client.post("/api/special-exams", headers=auth_header(student), json=valid(**overrides))


def test_options_lists_timetable_courses_and_defaults(client, student, timetable):
    res = client.get("/api/special-exams/options", headers=auth_header(student))
    assert res.status_code == 200
    body = res.get_json()
    assert body["batches"] == {"BCE08": ["Structural Analysis"]}
    assert body["defaults"]["id_number"] == "S123"


def test_student_registers_and_sees_pending(client, student, timetable, admin_team):
    res = submit(client, student)
    assert res.status_code == 201, res.get_json()
    assert res.get_json()["status"] == "pending"
    mine = client.get("/api/special-exams/mine", headers=auth_header(student)).get_json()
    assert [r["status"] for r in mine] == ["pending"]
    # Reviewers get an in-app heads-up, no email.
    assert Notification.query.filter_by(user_id=admin_team.id, type="special_exam_request").count() == 1
    assert EmailOutbox.query.count() == 0


@pytest.mark.parametrize("field,value", [
    ("full_name", ""), ("reason", "  "), ("phone", ""), ("id_number", ""), ("exam_date", ""),
    ("shift", "3"), ("phone", "call me"), ("batch", "XYZ01"), ("course_name", "Not a course"),
    ("exam_date", (date.today() + timedelta(days=1)).isoformat()),
])
def test_validation(client, student, timetable, field, value):
    assert submit(client, student, **{field: value}).status_code == 400


def test_duplicate_pending_rejected(client, student, timetable):
    assert submit(client, student).status_code == 201
    assert submit(client, student).status_code == 409


def test_decision_emails_and_shows_to_student(client, student, timetable, admin_team):
    req_id = submit(client, student).get_json()["id"]
    res = client.post(f"/api/special-exams/{req_id}/decision", headers=auth_header(admin_team),
                      json={"decision": "approved", "note": "Sit it on Saturday."})
    assert res.status_code == 200
    assert res.get_json()["decided_by_name"] == admin_team.name

    mine = client.get("/api/special-exams/mine", headers=auth_header(student)).get_json()[0]
    assert mine["status"] == "approved" and mine["decision_note"] == "Sit it on Saturday."
    email = EmailOutbox.query.filter_by(user_id=student.id).one()
    assert email.subject == "Special Exam Request Approved – Structural Analysis"
    assert "Sit it on Saturday." in email.html


def test_decline_without_note(client, student, timetable, admin_team):
    req_id = submit(client, student).get_json()["id"]
    res = client.post(f"/api/special-exams/{req_id}/decision", headers=auth_header(admin_team),
                      json={"decision": "declined"})
    assert res.status_code == 200 and res.get_json()["status"] == "declined"
    assert "Declined" in EmailOutbox.query.one().subject


def test_access_rules(client, student, lecturer, admin, timetable):
    req_id = submit(client, student).get_json()["id"]
    # Students and ordinary lecturers can't review.
    assert client.get("/api/special-exams", headers=auth_header(student)).status_code == 403
    assert client.get("/api/special-exams", headers=auth_header(lecturer)).status_code == 403
    assert client.post(f"/api/special-exams/{req_id}/decision", headers=auth_header(lecturer),
                       json={"decision": "approved"}).status_code == 403
    # Super Admin holds special_exams:* through RBAC.
    assert client.get("/api/special-exams", headers=auth_header(admin)).status_code == 200
    me = client.get("/api/auth/me", headers=auth_header(admin)).get_json()
    caps = me.get("capabilities") or me.get("user", {}).get("capabilities")
    assert caps["can_review_special_exams"] and caps["can_decide_special_exams"]


def test_excel_export_with_status_filter(client, student, timetable, admin_team):
    first = submit(client, student).get_json()["id"]
    submit(client, student, course_name="Structural Analysis", exam_date=date.today().isoformat(),
           reason="=HYPERLINK(\"http://evil\")")
    client.post(f"/api/special-exams/{first}/decision", headers=auth_header(admin_team), json={"decision": "approved"})

    res = client.get("/api/special-exams/export.xlsx", headers=auth_header(admin_team))
    assert res.status_code == 200
    ws = load_workbook(io.BytesIO(res.data)).active
    assert ws["A1"].value == "No." and ws.max_row == 3
    statuses = {ws.cell(row=r, column=13).value for r in (2, 3)}
    assert statuses == {"Approved", "Pending"}
    assert all(not str(ws.cell(row=r, column=11).value).startswith("=") for r in (2, 3))

    approved = client.get("/api/special-exams/export.xlsx?status=approved", headers=auth_header(admin_team))
    assert load_workbook(io.BytesIO(approved.data)).active.max_row == 2
    assert SpecialExamRequest.query.count() == 2


# ---------- Registration period ----------

@pytest.mark.closed_registration
def test_registration_closed_without_a_period(client, student, timetable):
    res = submit(client, student)
    assert res.status_code == 403 and "closed" in res.get_json()["error"]
    assert client.get("/api/special-exams/options", headers=auth_header(student)).get_json()["period"] is None


@pytest.mark.closed_registration
def test_period_controls_registration(client, student, timetable, admin_team):
    now = datetime.now()
    fmt = "%Y-%m-%dT%H:%M"
    # Upcoming: still closed.
    res = client.post("/api/special-exams/period", headers=auth_header(admin_team), json={
        "title": "Midterm", "opens_at": (now + timedelta(days=1)).strftime(fmt),
        "closes_at": (now + timedelta(days=3)).strftime(fmt),
    })
    assert res.status_code == 201 and res.get_json()["status"] == "upcoming"
    period_id = res.get_json()["id"]
    assert submit(client, student).status_code == 403
    options = client.get("/api/special-exams/options", headers=auth_header(student)).get_json()
    assert options["period"]["status"] == "upcoming"

    # Moved to start now: open.
    res = client.put(f"/api/special-exams/period/{period_id}", headers=auth_header(admin_team), json={
        "title": "Midterm", "opens_at": (now - timedelta(minutes=5)).strftime(fmt),
        "closes_at": (now + timedelta(days=3)).strftime(fmt),
    })
    assert res.get_json()["is_open"]
    created = submit(client, student)
    assert created.status_code == 201 and SpecialExamRequest.query.one().period_id == period_id

    # Closed early: registration stops.
    assert client.post(f"/api/special-exams/period/{period_id}/close",
                       headers=auth_header(admin_team)).get_json()["status"] == "closed"
    assert submit(client, student, exam_date=date.today().isoformat()).status_code == 403


@pytest.mark.closed_registration
def test_period_validation_and_access(client, student, lecturer, admin_team):
    now = datetime.now()
    body = {"opens_at": now.strftime("%Y-%m-%dT%H:%M"), "closes_at": (now - timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")}
    assert client.post("/api/special-exams/period", headers=auth_header(admin_team), json=body).status_code == 400
    assert client.post("/api/special-exams/period", headers=auth_header(admin_team), json={}).status_code == 400
    ok = {"opens_at": now.date().isoformat(), "closes_at": (now.date() + timedelta(days=2)).isoformat()}
    assert client.post("/api/special-exams/period", headers=auth_header(lecturer), json=ok).status_code == 403
    assert client.post("/api/special-exams/period", headers=auth_header(student), json=ok).status_code == 403
    assert client.post("/api/special-exams/period", headers=auth_header(admin_team), json=ok).status_code == 201
    # Overlapping periods are refused.
    assert client.post("/api/special-exams/period", headers=auth_header(admin_team), json=ok).status_code == 409
    current = client.get("/api/special-exams/period", headers=auth_header(admin_team)).get_json()["current"]
    assert current["is_open"]


# ---------- Evidence documents ----------

def submit_with_files(client, student, files, **overrides):
    data = valid(**overrides)
    data["documents"] = files
    return client.post("/api/special-exams", headers=auth_header(student), data=data,
                       content_type="multipart/form-data")


def test_documents_are_attached_and_protected(client, student, timetable, admin_team, lecturer):
    res = submit_with_files(client, student, [(io.BytesIO(PDF), "medical.pdf"),
                                              (io.BytesIO(b"\x89PNG\r\n\x1a\nxxxx"), "letter.png")])
    assert res.status_code == 201, res.get_json()
    body = res.get_json()
    assert [d["file_name"] for d in body["documents"]] == ["medical.pdf", "letter.png"]
    url = body["documents"][0]["url"].removeprefix("/api")

    assert client.get(f"/api{url}", headers=auth_header(student)).data == PDF
    assert client.get(f"/api{url}", headers=auth_header(admin_team)).status_code == 200
    assert client.get(f"/api{url}", headers=auth_header(lecturer)).status_code == 403
    other = make_user("other@example.com", role="student", batch="BCE08")
    assert client.get(f"/api{url}", headers=auth_header(other)).status_code == 403

    ws = load_workbook(io.BytesIO(client.get("/api/special-exams/export.xlsx", headers=auth_header(admin_team)).data)).active
    assert ws.cell(row=1, column=12).value == "Documents"
    assert "medical.pdf" in ws.cell(row=2, column=12).value


def test_documents_are_optional_and_validated(client, student, timetable):
    assert submit_with_files(client, student, []).status_code == 201
    bad = submit_with_files(client, student, [(io.BytesIO(b"MZ\x90\x00evil"), "virus.exe")],
                            exam_date=date.today().isoformat())
    assert bad.status_code == 400 and "virus.exe" in bad.get_json()["error"]
    many = [(io.BytesIO(PDF), f"f{i}.pdf") for i in range(6)]
    assert submit_with_files(client, student, many, exam_date=date.today().isoformat()).status_code == 400
    assert SpecialExamRequest.query.count() == 1


def test_notification_types_fit_the_column():
    """Postgres rejects a notification type longer than the column (SQLite in
    the tests doesn't), which once made approving a request fail with a 500."""
    import pathlib
    import re

    from app.models import Notification

    limit = Notification.__table__.c.type.type.length
    source = "\n".join(p.read_text() for p in pathlib.Path("app").rglob("*.py"))
    too_long = [t for t in set(re.findall(r'type="([a-z_]+)"', source)) if len(t) > limit]
    assert not too_long
