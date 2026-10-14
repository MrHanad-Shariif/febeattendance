"""Branded PDF/Excel downloads: every report endpoint gives both formats from
the same rows, named FEBE_<Report>_<date>.<ext>, behind its usual permission."""
import io
from datetime import date, datetime, time, timedelta

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models import Attendance, Timetable
from app.utils import report_export as rx
from tests.conftest import auth_header, make_user

ENDPOINTS = {
    "lecturer_attendance": ("/api/admin/attendance/report", "Lecturer_Attendance"),
    "timetable": ("/api/admin/timetable/export", "Timetable"),
    "absences": ("/api/admin/lecturer-absences/export", "Lecturer_Absences"),
    "course_completion": ("/api/course-completion/export", "Course_Completion"),
    "special_exams": ("/api/special-exams/export", "Special_Exam_Requests"),
}


@pytest.fixture()
def data(app, admin, lecturer):
    lecturer.name = "Abdirahman Mohamed Abdullahi Hassan Nur"  # long names must wrap, not clip
    t = Timetable(lecturer_id=lecturer.id, course_name="Structural Analysis and Design of Reinforced Concrete",
                  batch="BCE08", room="R101", days="Sat,Mon,Wed", start_time=time(7, 30), end_time=time(9, 0),
                  semester_total_sessions=34)
    db.session.add(t)
    d = date.today()
    db.session.add(Attendance(lecturer_id=lecturer.id, date=d, scheduled_start=datetime.combine(d, time(7, 30)),
                              status="late", checkin_at=datetime.combine(d, time(7, 45))))
    db.session.commit()
    return {"admin": admin, "lecturer": lecturer, "t": t}


@pytest.mark.parametrize("key", ENDPOINTS)
@pytest.mark.parametrize("fmt", ["pdf", "xlsx"])
def test_every_report_downloads_in_both_formats(client, data, key, fmt):
    url, slug = ENDPOINTS[key]
    resp = client.get(url, query_string={"format": fmt}, headers=auth_header(data["admin"]))
    assert resp.status_code == 200, (key, resp.data[:300])
    disposition = resp.headers["Content-Disposition"]
    assert f"FEBE_{slug}_{date.today().isoformat()}.{fmt}" in disposition
    if fmt == "pdf":
        assert resp.data.startswith(b"%PDF") and resp.mimetype == "application/pdf"
    else:
        ws = load_workbook(io.BytesIO(resp.data)).worksheets[0]
        assert ws["A1"].value == "SIMAD University"
        assert ws.freeze_panes is not None and ws.print_title_rows


def test_lecturer_attendance_report_has_late_minutes(client, data):
    resp = client.get("/api/admin/attendance/report?format=xlsx", headers=auth_header(data["admin"]))
    ws = load_workbook(io.BytesIO(resp.data)).worksheets[0]
    header = next(r for r in ws.iter_rows() if r[0].value == "Date")
    cols = [c.value for c in header]
    row = [c.value for c in ws[header[0].row + 1]]
    assert row[cols.index("Late min.")] == 15 and row[cols.index("Status")] == "Late"


def test_timetable_export_lists_each_meeting_day(client, data):
    resp = client.get("/api/admin/timetable/export?format=xlsx&batch=BCE08", headers=auth_header(data["admin"]))
    ws = load_workbook(io.BytesIO(resp.data)).worksheets[0]
    days = [r[0].value for r in ws.iter_rows() if r[0].value in ("Sat", "Mon", "Wed")]
    assert days == ["Sat", "Mon", "Wed"]
    assert client.get("/api/admin/timetable/export?format=xlsx&day=Funday", headers=auth_header(data["admin"])).status_code == 400


@pytest.mark.parametrize("key", ["lecturer_attendance", "timetable", "absences", "course_completion"])
def test_exports_keep_their_permissions(client, app, key):
    plain = make_user("plain@example.com", role="lecturer")
    student = make_user("stu@example.com", role="student")
    url = ENDPOINTS[key][0] + "?format=pdf"
    assert client.get(url, headers=auth_header(student)).status_code == 403
    if key != "course_completion":  # lecturers may see course completion
        assert client.get(url, headers=auth_header(plain)).status_code == 403


def test_excel_neutralises_formulas_and_keeps_numbers(app):
    rows = [{"name": "=HYPERLINK(\"http://evil\")", "n": 3, "p": 87.5}]
    report = rx.Report("T", "T", "P", [rx.Table("S", "T", [rx.Column("name", "Name"), rx.Column("n", "N", 5, "int"),
                                                         rx.Column("p", "P", 5, "pct")], rows)])
    ws = load_workbook(io.BytesIO(rx.to_xlsx(report))).worksheets[0]
    header = next(r for r in ws.iter_rows() if r[0].value == "Name")
    name, n, p = ws[header[0].row + 1][:3]
    assert name.value.startswith("'=") and n.value == 3 and p.value == 87.5 and p.number_format == '0.0"%"'


def test_large_pdf_is_split_into_pages_and_capped(app):
    rows = [{"name": f"Student {i}", "n": i} for i in range(rx.PDF_ROW_LIMIT + 50)]
    report = rx.Report("Big", "Big", "P", [rx.Table("S", "Big", [rx.Column("name", "Name", 30), rx.Column("n", "N", 6, "int")], rows)])
    pdf = rx.to_pdf(report)
    assert pdf.count(b"/Type /Page\n") > 30 or pdf.count(b"/Type /Page") > 30


def test_empty_reports_still_render(app):
    report = rx.Report("Empty", "Empty", "P", [rx.Table("S", "Empty", [rx.Column("a", "A")], [])])
    assert rx.to_pdf(report).startswith(b"%PDF")
    ws = load_workbook(io.BytesIO(rx.to_xlsx(report))).worksheets[0]
    assert any(c.value == "No records for this period and filter." for row in ws.iter_rows() for c in row)


def test_orientation_follows_the_columns():
    narrow = rx.Report("N", "N", "P", [rx.Table("S", "N", [rx.Column("a", "A", 20), rx.Column("b", "B", 20)], [])])
    wide = rx.Report("W", "W", "P", [rx.Table("S", "W", [rx.Column(str(i), str(i), 12) for i in range(10)], [])])
    assert rx._orientation(narrow) == "portrait" and rx._orientation(wide) == "landscape"


def test_branding_comes_from_settings(client, admin):
    client.put("/api/admin/settings", json={"report_university_name": "SIMAD University (Test)"}, headers=auth_header(admin))
    resp = client.get("/api/admin/timetable/export?format=xlsx", headers=auth_header(admin))
    assert load_workbook(io.BytesIO(resp.data)).worksheets[0]["A1"].value == "SIMAD University (Test)"
    assert client.put("/api/admin/settings", json={"report_footer": ""}, headers=auth_header(admin)).status_code == 400
    assert client.put("/api/admin/settings", json={"academic_year_start": "13-40"}, headers=auth_header(admin)).status_code == 400


@pytest.mark.parametrize("fmt", ["pdf", "xlsx"])
def test_student_report_download_for_admin_and_the_student_only(client, data, fmt):
    student = make_user("s1@example.com", role="student", student_id_number="C1190001", batch="BCE08")
    other = make_user("s2@example.com", role="student", student_id_number="C1190002", batch="BCE08")
    resp = client.get(f"/api/admin/students/{student.id}/report/export?format={fmt}", headers=auth_header(data["admin"]))
    assert resp.status_code == 200 and "FEBE_Student_Report_C1190001_" in resp.headers["Content-Disposition"]
    own = client.get(f"/api/student/report/export?format={fmt}", headers=auth_header(student))
    assert own.status_code == 200
    # A student can only ever download their own report.
    assert client.get(f"/api/admin/students/{student.id}/report/export?format={fmt}", headers=auth_header(other)).status_code == 403
    if fmt == "xlsx":
        wb = load_workbook(io.BytesIO(own.data))
        assert wb.sheetnames == ["Summary", "Sessions"]
        assert any("C1190001" in str(c.value) for row in wb["Summary"].iter_rows(max_row=8) for c in row if c.value)
