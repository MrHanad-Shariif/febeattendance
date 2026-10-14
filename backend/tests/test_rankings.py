"""Flexible rankings: any count or all, best/worst, criteria, per-class
student ranks, fair ties, filters, the new periods and PDF/Excel export."""
import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.utils.attendance_report import resolve_period
from app.utils.rankings import rank
from tests.conftest import auth_header
from tests.test_absences_and_reports import TODAY, _history, world  # noqa: F401 (fixture)

URL = "/api/admin/attendance-reports/ranking-table"


def _q(**extra):
    base = {"period": "range", "from": (TODAY - timedelta(days=14)).isoformat(), "to": (TODAY - timedelta(days=1)).isoformat()}
    return {**base, **extra}


def _get(client, world, **params):
    resp = client.get(URL, query_string=_q(**params), headers=auth_header(world["admin"]))
    assert resp.status_code == 200, resp.json
    return resp.json


def test_lecturers_best_first_with_late_minutes_as_tie_breaker(client, world):
    _history(world)  # A always on time; B 30 minutes late every day
    data = _get(client, world, audience="lecturers", metric="attendance_rate", limit="all")
    names = [(r["rank"], r["name"]) for r in data["rows"]]
    assert names == [(1, "Amina Lecturer"), (2, "Bashir Lecturer")]  # same attendance, fewer late minutes wins
    b = data["rows"][1]
    assert b["late_minutes"] == 14 * 30 and b["late"] == 14 and b["attendance_rate"] == 100.0
    assert "late minutes" in data["method"]


def test_worst_first_by_late_minutes(client, world):
    _history(world)
    data = _get(client, world, audience="lecturers", metric="late_minutes", order="worst")
    assert data["rows"][0]["name"] == "Bashir Lecturer" and data["rows"][0]["rank"] == 1
    assert data["summary"]["late_minutes"] == 420


@pytest.mark.parametrize("limit,expected", [("1", 1), ("2", 2), ("100", 2), ("all", 2), ("", 2)])
def test_any_number_of_results(client, world, limit, expected):
    _history(world)
    assert len(_get(client, world, audience="lecturers", limit=limit)["rows"]) == expected


@pytest.mark.parametrize("limit", ["0", "-3", "ten"])
def test_bad_limits_are_refused(client, world, limit):
    resp = client.get(URL, query_string=_q(audience="lecturers", limit=limit), headers=auth_header(world["admin"]))
    assert resp.status_code == 400


def test_unknown_metric_and_order_are_refused(client, world):
    h = auth_header(world["admin"])
    assert client.get(URL, query_string=_q(audience="students", metric="height"), headers=h).status_code == 400
    assert client.get(URL, query_string=_q(audience="students", order="random"), headers=h).status_code == 400


def test_students_ranked_within_their_class_and_excused_absences_do_not_count(client, world):
    _history(world)  # s1 always attends; s2 absent 2 of 3 days, excused the third
    data = _get(client, world, audience="students")
    assert data["per_class"] is True
    group = next(g for g in data["groups"] if g["class_name"] == "BCE08")
    s1, s2 = group["rows"]
    assert (s1["name"], s1["rank"], s1["attendance_rate"]) == (world["s1"].name, 1, 100.0)
    assert s2["rank"] == 2 and s2["justified_absences"] == 4
    assert s2["accountable_sessions"] == 14 - 4 and s2["attendance_rate"] == 0.0
    worst = _get(client, world, audience="students", order="worst")
    assert worst["groups"][0]["rows"][0]["name"] == world["s2"].name


def test_minimum_sessions_excludes_people_with_too_little_data(client, world):
    _history(world)
    data = _get(client, world, audience="students", min_sessions="11")
    assert data["excluded_count"] == 1 and data["eligible_count"] == 1


def test_department_and_lecturer_filters(client, world):
    _history(world)
    world["a"].department = "Civil Engineering"
    db.session.commit()
    data = _get(client, world, audience="lecturers", department="Civil Engineering")
    assert [r["name"] for r in data["rows"]] == ["Amina Lecturer"]
    data = _get(client, world, audience="lecturers", lecturer_id=str(world["b"].id))
    assert [r["name"] for r in data["rows"]] == ["Bashir Lecturer"]


def test_ties_share_a_rank_and_cut_off_keeps_ties():
    people = [{"name": n, "x": v} for n, v in [("Cali", 90), ("Asha", 90), ("Bile", 80), ("Dahir", 70)]]
    out = rank(people, [("x", True)])
    assert [(p["rank"], p["name"], p["tied"]) for p in out] == [
        (1, "Asha", True), (1, "Cali", True), (3, "Bile", False), (4, "Dahir", False)]
    assert [p["name"] for p in rank(people, [("x", True)], limit=1)] == ["Asha", "Cali"]
    worst = rank(people, [("x", True)], order="worst")
    assert [(p["rank"], p["name"]) for p in worst] == [(1, "Dahir"), (2, "Bile"), (3, "Asha"), (3, "Cali")]


def test_missing_rate_sorts_last():
    people = [{"name": "A", "x": None}, {"name": "B", "x": 10}]
    assert [p["name"] for p in rank(people, [("x", True)])] == ["B", "A"]


def test_new_periods(app):
    d = date(2026, 10, 14)  # a Wednesday
    week = resolve_period({"period": "weekly", "date": d.isoformat()}, today=d)
    assert (week["from"], week["to"]) == (date(2026, 10, 12), date(2026, 10, 18))
    year = resolve_period({"period": "academic_year"}, today=d)
    assert (year["from"], year["to"]) == (date(2026, 9, 1), date(2027, 8, 31))
    assert resolve_period({"period": "academic_year", "year": "2025"}, today=d)["from"] == date(2025, 9, 1)
    assert resolve_period({"period": "all"}, today=d)["to"] == d


@pytest.mark.parametrize("fmt", ["pdf", "xlsx"])
def test_ranking_export(client, world, fmt):
    _history(world)
    resp = client.get("/api/admin/attendance-reports/export",
                      query_string=_q(audience="ranking", ranking_audience="students", format=fmt, limit="10"),
                      headers=auth_header(world["admin"]))
    assert resp.status_code == 200, resp.data[:300]
    assert "FEBE_Student_Rankings_" in resp.headers["Content-Disposition"]
    if fmt == "xlsx":
        wb = load_workbook(io.BytesIO(resp.data))
        assert "Class BCE08" in wb.sheetnames
        ws = wb["Class BCE08"]
        header_row = next(r for r in ws.iter_rows() if r[0].value == "Rank")
        pct_cell = ws.cell(row=header_row[0].row + 1, column=10)
        assert isinstance(pct_cell.value, (int, float)) and pct_cell.number_format == '0.0"%"'
    else:
        assert resp.data.startswith(b"%PDF")


def test_rankings_need_report_permission(client, world):
    resp = client.get(URL, query_string=_q(audience="students"), headers=auth_header(world["s1"]))
    assert resp.status_code == 403
