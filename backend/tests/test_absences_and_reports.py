"""Lecturer absences (exemptions, review, close/cancel, conflicts), the
session-level daily/monthly/semester reports, rankings, Excel exports and
student class correction."""
import io
from datetime import date, datetime, time, timedelta

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models import (
    AbsenceReviewItem, Attendance, LecturerAbsence, Role, Setting, StudentAttendance, StudentClassChange, Timetable,
)
from app.utils.absences import AbsenceIndex
from app.utils.schedule import day_abbr
from app.utils.student_status import compute_course_stats
from tests.conftest import auth_header, make_user

TODAY = date.today()


def _past(weekday_offset_weeks=1, days_back=7):
    return TODAY - timedelta(days=days_back)


def _set(key, value):
    row = Setting.query.get(key)
    if row:
        row.value = value
    else:
        db.session.add(Setting(key=key, value=value))
    db.session.commit()


@pytest.fixture()
def world(app, admin):
    """Two lecturers. Lecturer A teaches BCE08 Statics (08:00-10:00) and
    BCE09 Hydraulics (10:30-12:00) every day of the week, so any date works.
    Lecturer B teaches BCE08 Surveying (13:00-15:00)."""
    every_day = "Sat,Sun,Mon,Tue,Wed,Thu,Fri"
    a = make_user("lect.a@example.com")
    a.name = "Amina Lecturer"
    b = make_user("lect.b@example.com")
    b.name = "Bashir Lecturer"
    t1 = Timetable(lecturer_id=a.id, course_name="Statics", batch="BCE08", days=every_day,
                   start_time=time(8, 0), end_time=time(10, 0))
    t2 = Timetable(lecturer_id=a.id, course_name="Hydraulics", batch="BCE09", days=every_day,
                   start_time=time(10, 30), end_time=time(12, 0))
    t3 = Timetable(lecturer_id=b.id, course_name="Surveying", batch="BCE08", days=every_day,
                   start_time=time(13, 0), end_time=time(15, 0))
    db.session.add_all([t1, t2, t3])
    s1 = make_user("s1@example.com", role="student", student_id_number="S1", batch="BCE08")
    s2 = make_user("s2@example.com", role="student", student_id_number="S2", batch="BCE08")
    s9 = make_user("s9@example.com", role="student", student_id_number="S9", batch="BCE09")
    for s in (s1, s2, s9):
        s.created_at = datetime.now() - timedelta(days=200)
    _set("semester_start_date", (TODAY - timedelta(days=60)).isoformat())
    _set("semester_end_date", (TODAY + timedelta(days=60)).isoformat())
    db.session.commit()
    return dict(admin=admin, a=a, b=b, t1=t1, t2=t2, t3=t3, s1=s1, s2=s2, s9=s9)


def _absence_body(w, **extra):
    body = {
        "lecturer_id": w["a"].id,
        "start_date": (TODAY - timedelta(days=3)).isoformat(),
        "end_date": (TODAY - timedelta(days=1)).isoformat(),
        "reason": "medical_leave",
        "remarks": "Hospital note on file",
    }
    body.update(extra)
    return body


def _sa(student, entry, d, status="absent", checkin=None):
    r = StudentAttendance(student_id=student.id, timetable_id=entry.id, date=d,
                          scheduled_start=datetime.combine(d, entry.start_time), status=status, checkin_at=checkin)
    db.session.add(r)
    db.session.commit()
    return r


# ---------- recording an absence ----------

def test_preview_lists_affected_classes_and_counts(client, world):
    body = _absence_body(world)
    resp = client.post("/api/admin/lecturer-absences/preview", json=body, headers=auth_header(world["admin"]))
    assert resp.status_code == 200, resp.json
    data = resp.json
    assert data["total_sessions"] == 6  # 3 days x 2 classes
    assert {c["batch"] for c in data["classes"]} == {"BCE08", "BCE09"}
    assert data["conflicts"] == []
    # Nothing is saved by a preview.
    assert LecturerAbsence.query.count() == 0


def test_scoped_absence_only_covers_chosen_class_and_days(client, world):
    d = TODAY - timedelta(days=2)
    body = _absence_body(world, batches=["BCE09"], days=[day_abbr(d)])
    data = client.post("/api/admin/lecturer-absences/preview", json=body, headers=auth_header(world["admin"])).json
    assert data["total_sessions"] == 1
    assert data["classes"][0]["course_name"] == "Hydraulics"


def test_create_requires_reason_and_valid_dates(client, world):
    h = auth_header(world["admin"])
    assert client.post("/api/admin/lecturer-absences", json=_absence_body(world, reason="nap"), headers=h).status_code == 400
    assert client.post("/api/admin/lecturer-absences", json=_absence_body(world, end_date=(TODAY - timedelta(days=9)).isoformat()), headers=h).status_code == 400
    assert client.post("/api/admin/lecturer-absences", json=_absence_body(world, end_date=None), headers=h).status_code == 400
    assert client.post("/api/admin/lecturer-absences", json=_absence_body(world, reason="other"), headers=h).status_code == 400
    ok = client.post("/api/admin/lecturer-absences", json=_absence_body(world, reason="other", custom_reason="Conference"), headers=h)
    assert ok.status_code == 201
    assert ok.json["reason_label"] == "Conference"
    assert ok.json["created_by_name"] == world["admin"].name


def test_duplicate_or_overlapping_absence_is_refused(client, world):
    h = auth_header(world["admin"])
    assert client.post("/api/admin/lecturer-absences", json=_absence_body(world), headers=h).status_code == 201
    dup = client.post("/api/admin/lecturer-absences", json=_absence_body(world, start_date=(TODAY - timedelta(days=2)).isoformat()), headers=h)
    assert dup.status_code == 409
    # A different lecturer on the same dates is fine.
    assert client.post("/api/admin/lecturer-absences", json=_absence_body(world, lecturer_id=world["b"].id), headers=h).status_code == 201
    # Same lecturer, non-overlapping class scope is fine too.
    first = LecturerAbsence.query.filter_by(lecturer_id=world["a"].id).first()
    first.batches = ["BCE08"]
    db.session.commit()
    AbsenceIndex.invalidate()
    assert client.post("/api/admin/lecturer-absences", json=_absence_body(world, batches=["BCE09"]), headers=h).status_code == 201


# ---------- exemptions ----------

def test_past_date_absence_flags_existing_records_and_exempts_them(client, world):
    d = TODAY - timedelta(days=2)
    # Recorded before anyone knew the lecturer was away: auto-absent rows.
    r1 = _sa(world["s1"], world["t1"], d, "absent")
    r2 = _sa(world["s2"], world["t1"], d, "absent")
    stats_before = compute_course_stats(world["s1"].id, "BCE08", "Statics")
    assert stats_before["absent"] == 1

    resp = client.post("/api/admin/lecturer-absences", json=_absence_body(world), headers=auth_header(world["admin"]))
    assert resp.status_code == 201
    assert resp.json["flagged_for_review"] == 2

    # Rows are untouched...
    assert db.session.get(StudentAttendance, r1.id).status == "absent"
    # ...but no longer counted.
    AbsenceIndex.invalidate()
    stats = compute_course_stats(world["s1"].id, "BCE08", "Statics")
    assert stats["absent"] == 0
    assert stats["exempted_records"] == 1
    assert stats["exempted_sessions"] == 3
    assert stats["total_sessions"] == stats_before["total_sessions"] - 3

    # Admin keeps one row as recorded: it counts again.
    absence_id = resp.json["id"]
    item = AbsenceReviewItem.query.filter_by(kind="student", record_id=r2.id).one()
    kept = client.post(f"/api/admin/lecturer-absences/{absence_id}/review",
                       json={"item_ids": [item.id], "resolution": "kept"}, headers=auth_header(world["admin"]))
    assert kept.status_code == 200
    AbsenceIndex.invalidate()
    assert compute_course_stats(world["s2"].id, "BCE08", "Statics")["absent"] == 1
    assert compute_course_stats(world["s1"].id, "BCE08", "Statics")["absent"] == 0
    assert any(e["action"] == "review_resolved" for e in kept.json["events"])


def test_unaffected_lecturer_and_class_still_count(client, world):
    d = TODAY - timedelta(days=2)
    _sa(world["s1"], world["t3"], d, "absent")  # lecturer B's class
    client.post("/api/admin/lecturer-absences", json=_absence_body(world), headers=auth_header(world["admin"]))
    AbsenceIndex.invalidate()
    assert compute_course_stats(world["s1"].id, "BCE08", "Surveying")["absent"] == 1


def test_job_does_not_mark_students_absent_for_exempt_session(app, client, world):
    from app.jobs import mark_student_absentees

    now = datetime.now()
    if now.hour < 1:
        pytest.skip("needs a class that started over 20 minutes ago today")
    start = (now - timedelta(hours=1)).time().replace(second=0, microsecond=0)
    world["t1"].start_time = start
    world["t1"].end_time = None
    world["t3"].start_time = start
    world["t3"].end_time = None
    db.session.commit()
    client.post("/api/admin/lecturer-absences", json=_absence_body(world, start_date=TODAY.isoformat(),
                                                                    end_date=TODAY.isoformat()),
                headers=auth_header(world["admin"]))
    mark_student_absentees(app)
    # Lecturer A's class: exempted, nobody marked absent. Lecturer B's: normal.
    assert StudentAttendance.query.filter_by(timetable_id=world["t1"].id, date=TODAY).count() == 0
    assert StudentAttendance.query.filter_by(timetable_id=world["t3"].id, date=TODAY, status="absent").count() == 2


def test_student_today_shows_exempted_and_checkin_is_refused(app, client, world):
    from app.student.routes import CheckinRefused, _open_session
    from app.utils.settings import get_all_settings

    client.post("/api/admin/lecturer-absences", json=_absence_body(world, start_date=TODAY.isoformat(),
                                                                    end_date=TODAY.isoformat(), batches=["BCE08"]),
                headers=auth_header(world["admin"]))
    rows = client.get("/api/student/today", headers=auth_header(world["s1"])).json
    statics = next(r for r in rows if r["course_name"] == "Statics")
    surveying = next(r for r in rows if r["course_name"] == "Surveying")
    assert statics["status"] == "exempted" and statics["status_label"] == "Class Session Exempted"
    assert surveying["exempted"] is False
    with pytest.raises(CheckinRefused) as err:
        _open_session(world["s1"], world["t1"].id, TODAY, get_all_settings())
    assert "exempted" in err.value.message.lower()


def test_ongoing_absence_then_early_close_resumes_attendance(client, world):
    h = auth_header(world["admin"])
    resp = client.post("/api/admin/lecturer-absences",
                       json=_absence_body(world, start_date=(TODAY - timedelta(days=5)).isoformat(), end_date=None, ongoing=True),
                       headers=h)
    assert resp.status_code == 201
    assert resp.json["ongoing"] is True and resp.json["status"] == "active"
    absence = LecturerAbsence.query.get(resp.json["id"])
    assert absence.covers_date(TODAY + timedelta(days=30))
    # While ongoing, only the sessions already passed (5 days back + today) come off the planned total.
    AbsenceIndex.invalidate()
    assert compute_course_stats(world["s1"].id, "BCE08", "Statics")["exempted_sessions"] == 6

    returned = TODAY - timedelta(days=1)
    closed = client.post(f"/api/admin/lecturer-absences/{absence.id}/close", json={"return_date": returned.isoformat()}, headers=h)
    assert closed.status_code == 200, closed.json
    assert closed.json["status"] == "completed"
    assert closed.json["actual_end_date"] == (returned - timedelta(days=1)).isoformat()
    AbsenceIndex.invalidate()
    index = AbsenceIndex()
    assert index.covering(world["t1"], returned - timedelta(days=1)) is not None
    assert index.covering(world["t1"], returned) is None
    assert index.covering(world["t1"], TODAY) is None
    # Can't close on or before the start date.
    assert client.post(f"/api/admin/lecturer-absences/{absence.id}/close",
                       json={"return_date": (TODAY - timedelta(days=5)).isoformat()}, headers=h).status_code == 400


def test_scheduled_status_and_edit_audit_trail(client, world):
    h = auth_header(world["admin"])
    resp = client.post("/api/admin/lecturer-absences",
                       json=_absence_body(world, start_date=(TODAY + timedelta(days=3)).isoformat(),
                                          end_date=(TODAY + timedelta(days=5)).isoformat(), reason="training"),
                       headers=h)
    assert resp.json["status"] == "scheduled"
    edited = client.put(f"/api/admin/lecturer-absences/{resp.json['id']}",
                        json={"end_date": (TODAY + timedelta(days=7)).isoformat(), "change_note": "Course extended"}, headers=h)
    assert edited.status_code == 200, edited.json
    actions = [e["action"] for e in edited.json["events"]]
    assert actions[0] == "edited" and "created" in actions
    assert edited.json["events"][0]["old_value"]["end_date"] == (TODAY + timedelta(days=5)).isoformat()


def test_cancelled_absence_no_longer_exempts(client, world):
    h = auth_header(world["admin"])
    d = TODAY - timedelta(days=2)
    _sa(world["s1"], world["t1"], d, "absent")
    resp = client.post("/api/admin/lecturer-absences", json=_absence_body(world), headers=h)
    assert client.post(f"/api/admin/lecturer-absences/{resp.json['id']}/cancel", json={}, headers=h).status_code == 400
    cancelled = client.post(f"/api/admin/lecturer-absences/{resp.json['id']}/cancel", json={"reason": "Entered by mistake"}, headers=h)
    assert cancelled.status_code == 200
    assert cancelled.json["status"] == "cancelled"
    AbsenceIndex.invalidate()
    assert compute_course_stats(world["s1"].id, "BCE08", "Statics")["absent"] == 1
    assert client.put(f"/api/admin/lecturer-absences/{resp.json['id']}", json={"remarks": "x"}, headers=h).status_code == 409
    # Still listed (never deleted) and filterable by status.
    listed = client.get("/api/admin/lecturer-absences?status=cancelled", headers=h).json
    assert [r["id"] for r in listed] == [resp.json["id"]]


def test_absence_list_filters(client, world):
    h = auth_header(world["admin"])
    client.post("/api/admin/lecturer-absences", json=_absence_body(world), headers=h)
    client.post("/api/admin/lecturer-absences", json=_absence_body(world, lecturer_id=world["b"].id, reason="annual_leave"), headers=h)
    assert len(client.get("/api/admin/lecturer-absences", headers=h).json) == 2
    assert len(client.get(f"/api/admin/lecturer-absences?lecturer_id={world['b'].id}", headers=h).json) == 1
    assert len(client.get("/api/admin/lecturer-absences?reason=annual_leave", headers=h).json) == 1
    assert len(client.get(f"/api/admin/lecturer-absences?from={(TODAY + timedelta(days=1)).isoformat()}", headers=h).json) == 0
    row = client.get(f"/api/admin/lecturer-absences?lecturer_id={world['a'].id}", headers=h).json[0]
    assert row["affected_sessions"] == 6
    assert row["affected_classes"] == ["BCE08", "BCE09"]


def test_absence_endpoints_are_admin_only(client, world):
    body = _absence_body(world)
    for user in (world["a"], world["s1"]):
        h = auth_header(user)
        assert client.get("/api/admin/lecturer-absences", headers=h).status_code in (401, 403)
        assert client.post("/api/admin/lecturer-absences", json=body, headers=h).status_code in (401, 403)
    # A role with view only can look but not record.
    viewer = make_user("viewer@example.com", role="admin")
    from app.models import Permission
    role = Role(name="Absence viewer", permissions=[Permission.query.filter_by(code="lecturer_absences:view").one()])
    viewer.roles.append(role)
    db.session.commit()
    h = auth_header(viewer)
    assert client.get("/api/admin/lecturer-absences", headers=h).status_code == 200
    assert client.post("/api/admin/lecturer-absences", json=body, headers=h).status_code == 403
    assert client.get("/api/admin/attendance-reports/lecturers", headers=h).status_code == 403


# ---------- lecturer daily report ----------

def test_daily_report_numbers_sessions_and_splits_justified(client, world):
    d = TODAY - timedelta(days=10)
    db.session.add(Attendance(lecturer_id=world["b"].id, date=d, scheduled_start=datetime.combine(d, time(13)),
                              status="absent"))
    db.session.add(Attendance(lecturer_id=world["a"].id, date=d, scheduled_start=datetime.combine(d, time(8)),
                              scheduled_end=datetime.combine(d, time(12)), status="late",
                              checkin_at=datetime.combine(d, time(8, 20)), checkout_at=datetime.combine(d, time(12, 5))))
    db.session.commit()
    rows = client.get(f"/api/admin/attendance/daily?date={d.isoformat()}", headers=auth_header(world["admin"])).json["rows"]
    a_rows = [r for r in rows if r["lecturer_id"] == world["a"].id]
    assert [r["session_number"] for r in a_rows] == [1, 2]
    assert [r["class_name"] for r in a_rows] == ["BCE08", "BCE09"]
    assert a_rows[0]["course_name"] == "Statics" and a_rows[0]["scheduled_time"] == "08:00–10:00"
    assert a_rows[0]["status"] == "late"          # 20 min after 08:00
    assert a_rows[1]["status"] == "on_time"       # before 10:30
    assert a_rows[0]["checkin_at"] == "08:20" and a_rows[0]["checkout_at"] == "12:05"
    b_row = next(r for r in rows if r["lecturer_id"] == world["b"].id)
    assert b_row["status_label"] == "Absent (unjustified)"
    assert b_row["category"] == "absent"
    assert b_row["remarks"] is None  # no reason recorded -> left blank


def test_daily_report_marks_lecturer_unavailable_as_justified(client, world):
    d = TODAY - timedelta(days=2)
    db.session.add(Attendance(lecturer_id=world["a"].id, date=d, scheduled_start=datetime.combine(d, time(8)), status="absent"))
    db.session.commit()
    client.post("/api/admin/lecturer-absences", json=_absence_body(world), headers=auth_header(world["admin"]))
    rows = client.get(f"/api/admin/attendance/daily?date={d.isoformat()}", headers=auth_header(world["admin"])).json["rows"]
    a_rows = [r for r in rows if r["lecturer_id"] == world["a"].id]
    assert all(r["status"] == "lecturer_unavailable" and r["category"] == "justified" for r in a_rows)
    assert a_rows[0]["justification"] == "Justified"
    assert "Medical leave" in a_rows[0]["remarks"]
    # The old monthly summary also counts the day as a justified absence.
    summary = client.get(f"/api/admin/summary?month={d:%Y-%m}", headers=auth_header(world["admin"])).json["rows"]
    a_sum = next(r for r in summary if r["lecturer_id"] == world["a"].id)
    assert a_sum["absent_justified"] == 1 and a_sum["absent_unjustified"] == 0


def test_lecturer_checkin_times_ignore_covered_classes(app, client, world):
    from app.utils.absences import split_day_classes
    from app.utils.schedule import day_bounds, get_day_classes

    client.post("/api/admin/lecturer-absences", json=_absence_body(world, start_date=TODAY.isoformat(),
                                                                    end_date=TODAY.isoformat(), batches=["BCE08"]),
                headers=auth_header(world["admin"]))
    entries = get_day_classes(world["a"].id, TODAY, set())
    teaching, covered = split_day_classes(entries, TODAY)
    assert [e.course_name for e in teaching] == ["Hydraulics"]
    first_start, _ = day_bounds(teaching, TODAY)
    assert first_start.time() == time(10, 30)  # the 08:00 class is covered


# ---------- reports ----------

def _history(world):
    """Two weeks of history: A always on time, B late every day, plus student rows."""
    for i in range(1, 15):
        d = TODAY - timedelta(days=i)
        db.session.add(Attendance(lecturer_id=world["a"].id, date=d, scheduled_start=datetime.combine(d, time(8)),
                                  status="on_time", checkin_at=datetime.combine(d, time(7, 55)),
                                  checkout_at=datetime.combine(d, time(12, 1))))
        db.session.add(Attendance(lecturer_id=world["b"].id, date=d, scheduled_start=datetime.combine(d, time(13)),
                                  status="late", checkin_at=datetime.combine(d, time(13, 30)),
                                  checkout_at=datetime.combine(d, time(15, 0))))
        db.session.add(StudentAttendance(student_id=world["s1"].id, timetable_id=world["t1"].id, date=d,
                                         scheduled_start=datetime.combine(d, time(8)),
                                         status="late" if i % 2 else "on_time", checkin_at=datetime.combine(d, time(8, 15))))
        db.session.add(StudentAttendance(student_id=world["s2"].id, timetable_id=world["t1"].id, date=d,
                                         scheduled_start=datetime.combine(d, time(8)),
                                         status="absent" if i % 3 else "present", remarks=None if i % 3 else "Sick note"))
    db.session.commit()


def test_lecturer_summary_and_detailed_reports(client, world):
    _history(world)
    h = auth_header(world["admin"])
    q = f"period=range&from={(TODAY - timedelta(days=14)).isoformat()}&to={(TODAY - timedelta(days=1)).isoformat()}"
    summary = client.get(f"/api/admin/attendance-reports/lecturers?{q}&view=summary", headers=h).json
    a = next(r for r in summary["rows"] if r["id"] == world["a"].id)
    b = next(r for r in summary["rows"] if r["id"] == world["b"].id)
    assert a["present"] == 28 and a["present_rate"] == 100.0
    assert b["late"] == 14 and b["late_rate"] == 100.0
    assert summary["totals"]["counted_sessions"] == 42

    detailed = client.get(f"/api/admin/attendance-reports/lecturers?{q}&view=detailed&category=late", headers=h).json
    assert detailed["total_rows"] == 14
    assert all(r["category"] == "late" for r in detailed["rows"])
    assert client.get(f"/api/admin/attendance-reports/lecturers?{q}&category=bogus", headers=h).status_code == 400


def test_student_report_excludes_exempted_sessions_from_percentages(client, world):
    _history(world)
    h = auth_header(world["admin"])
    q = f"period=range&from={(TODAY - timedelta(days=14)).isoformat()}&to={(TODAY - timedelta(days=1)).isoformat()}"
    before = client.get(f"/api/admin/attendance-reports/students?{q}&student_id={world['s2'].id}", headers=h).json
    s2_before = before["rows"][0]
    assert s2_before["counted_sessions"] == 14
    assert s2_before["justified"] == 4 and s2_before["absent"] == 10

    client.post("/api/admin/lecturer-absences", json=_absence_body(world), headers=h)  # covers days 1-3
    AbsenceIndex.invalidate()
    after = client.get(f"/api/admin/attendance-reports/students?{q}&student_id={world['s2'].id}", headers=h).json["rows"][0]
    assert after["counted_sessions"] == 11
    assert after["exempted"] == 3
    assert after["absent"] + after["justified"] == 11
    detailed = client.get(f"/api/admin/attendance-reports/students?{q}&view=detailed&category=exempted&batch=BCE08",
                          headers=h).json
    # s1 and s2 had stored rows (now exempted); nothing else in BCE08 is covered.
    assert detailed["total_rows"] == 6
    assert all(r["status_label"] == "Class Session Exempted" for r in detailed["rows"])
    # BCE09's student had no rows: exempted sessions are still listed for them.
    bce09 = client.get(f"/api/admin/attendance-reports/students?{q}&view=detailed&batch=BCE09", headers=h).json
    assert bce09["total_rows"] == 3 and {r["category"] for r in bce09["rows"]} == {"exempted"}


def test_monthly_and_semester_periods_resolve(client, world):
    h = auth_header(world["admin"])
    m = client.get(f"/api/admin/attendance-reports/lecturers?period=monthly&month={TODAY:%Y-%m}", headers=h).json
    assert m["from"] == TODAY.replace(day=1).isoformat()
    s = client.get("/api/admin/attendance-reports/students?period=semester", headers=h).json
    assert s["from"] == (TODAY - timedelta(days=60)).isoformat()
    assert client.get("/api/admin/attendance-reports/students?period=range&from=2026-01-01", headers=h).status_code == 400


def test_rankings_top_three_ties_and_insufficient_data(client, world):
    _history(world)
    world["t1"].semester_total_sessions = 30  # 10 absences of 30 planned = 33 %: over the 25 % limit
    c = make_user("lect.c@example.com")
    c.name = "Cali Lecturer"
    db.session.add(Timetable(lecturer_id=c.id, course_name="Drawing", batch="BAR01", days="Sat,Sun,Mon,Tue,Wed,Thu,Fri",
                             start_time=time(16), end_time=time(17)))
    d = TODAY - timedelta(days=1)
    db.session.add(Attendance(lecturer_id=c.id, date=d, scheduled_start=datetime.combine(d, time(16)), status="on_time",
                              checkin_at=datetime.combine(d, time(15, 55))))
    db.session.commit()
    q = f"period=range&from={(TODAY - timedelta(days=14)).isoformat()}&to={(TODAY - timedelta(days=1)).isoformat()}"
    data = client.get(f"/api/admin/attendance-reports/rankings?{q}&min_sessions=5", headers=auth_header(world["admin"])).json
    lr = data["lecturers"]
    assert lr["top"][0]["lecturer_name"] == "Amina Lecturer" and lr["top"][0]["rank"] == 1
    assert lr["most_late"][0]["lecturer_name"] == "Bashir Lecturer"
    # Cali taught only 1 recorded session but 14 were scheduled -> 13 unjustified absences, enough data.
    assert any(r["lecturer_name"] == "Cali Lecturer" for r in lr["most_unjustified_absences"])
    sr = data["students"]
    assert sr["repeated_lateness"][0]["student_name"] == world["s1"].name
    assert sr["frequent_absences"][0]["student_name"] == world["s2"].name
    s2_warning = next(w for w in data["warnings"]["rows"] if w["student_id"] == world["s2"].id)
    assert s2_warning["level"] == "blocked" and s2_warning["absence_percentage"] == 33.3


def test_ranking_ties_are_flagged():
    from app.utils.attendance_report import _rank

    items = [{"n": "a", "v": 5}, {"n": "b", "v": 5}, {"n": "c", "v": 3}, {"n": "d", "v": 3}, {"n": "e", "v": 1}]
    top = _rank(items, lambda i: -i["v"], limit=3)
    assert [(i["n"], i["rank"], i["tied"]) for i in top] == [("a", 1, True), ("b", 1, True), ("c", 3, True), ("d", 3, True)]


def test_lecturer_with_justified_absences_is_not_penalised():
    from app.utils.attendance_report import lecturer_rankings

    rows = ([{"lecturer_id": 1, "lecturer_name": "X", "category": "present"}] * 6
            + [{"lecturer_id": 1, "lecturer_name": "X", "category": "justified"}] * 6
            + [{"lecturer_id": 2, "lecturer_name": "Y", "category": "present"}] * 6
            + [{"lecturer_id": 2, "lecturer_name": "Y", "category": "absent"}] * 6)
    r = lecturer_rankings(rows, min_sessions=5)
    assert r["top"][0]["lecturer_name"] == "X" and r["top"][0]["on_time_rate"] == 100.0
    assert r["highest_absence_rate"][0]["lecturer_name"] == "Y"


@pytest.mark.parametrize("audience", ["lecturers", "students"])
@pytest.mark.parametrize("period", ["daily", "monthly", "semester"])
@pytest.mark.parametrize("view", ["summary", "detailed"])
def test_excel_exports(client, world, audience, period, view):
    _history(world)
    resp = client.get(f"/api/admin/attendance-reports/export?audience={audience}&period={period}&view={view}"
                      f"&date={(TODAY - timedelta(days=1)).isoformat()}",
                      headers=auth_header(world["admin"]))
    assert resp.status_code == 200, resp.data[:300]
    assert resp.mimetype.endswith("spreadsheetml.sheet")
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.worksheets[0]
    assert ws["A1"].value == "SIMAD University"
    assert "FEBEMS" in ws["A2"].value
    assert "Attendance Report" in ws["A4"].value
    assert ws["A5"].value.startswith("Reporting period:")
    assert resp.headers["Content-Disposition"].endswith(".xlsx")
    assert f"FEBE_{audience[:-1].title()}_Attendance_{view.title()}_" in resp.headers["Content-Disposition"]
    assert ws.page_setup.orientation == "landscape"
    assert ws.print_title_rows is not None
    if view == "detailed":
        assert ws.title == "Detailed" and "Summary" in wb.sheetnames


@pytest.mark.parametrize("audience", ["lecturers", "students"])
@pytest.mark.parametrize("view", ["summary", "detailed"])
def test_pdf_exports(client, world, audience, view):
    _history(world)
    resp = client.get(f"/api/admin/attendance-reports/export?audience={audience}&period=monthly&view={view}&format=pdf",
                      headers=auth_header(world["admin"]))
    assert resp.status_code == 200, resp.data[:300]
    assert resp.mimetype == "application/pdf" and resp.data.startswith(b"%PDF")
    assert resp.headers["Content-Disposition"].endswith(".pdf")


def test_export_rejects_unknown_format(client, world):
    resp = client.get("/api/admin/attendance-reports/export?audience=lecturers&format=docx", headers=auth_header(world["admin"]))
    assert resp.status_code == 400


def test_rankings_excel_export(client, world):
    _history(world)
    resp = client.get("/api/admin/attendance-reports/export?audience=rankings&period=semester", headers=auth_header(world["admin"]))
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.data))
    assert {"Top lecturers", "Most late", "Student warnings"} <= set(wb.sheetnames)


# ---------- student class correction ----------

def test_class_correction_keeps_history_and_is_audited(client, world):
    h = auth_header(world["admin"])
    d = TODAY - timedelta(days=3)
    r = _sa(world["s1"], world["t1"], d, "on_time", checkin=datetime.combine(d, time(8)))
    s1_id, s1_email = world["s1"].id, world["s1"].email

    resp = client.post(f"/api/admin/students/{s1_id}/class", json={"batch": "BCE09", "reason": "Registered in the wrong class"}, headers=h)
    assert resp.status_code == 200, resp.json
    assert resp.json["batch"] == "BCE09" and resp.json["old_batch"] == "BCE08"
    assert resp.json["id"] == s1_id and resp.json["email"] == s1_email  # same account

    change = StudentClassChange.query.filter_by(student_id=s1_id).one()
    assert (change.old_batch, change.new_batch, change.changed_by_name) == ("BCE08", "BCE09", world["admin"].name)
    assert db.session.get(StudentAttendance, r.id) is not None  # history kept

    report = client.get(f"/api/admin/students/{s1_id}/report", headers=h).json
    assert {c["course_name"] for c in report["summary"]} >= {"Statics", "Hydraulics"}  # old class history + new class

    today_rows = client.get("/api/student/today", headers=auth_header(world["s1"])).json
    assert {row["batch"] for row in today_rows} == {"BCE09"}  # future attendance follows the new class

    history = client.get(f"/api/admin/students/{s1_id}/class-history", headers=h).json
    assert history[0]["reason"] == "Registered in the wrong class"


def test_class_correction_validation_and_permissions(client, world):
    h = auth_header(world["admin"])
    sid = world["s1"].id
    assert client.post(f"/api/admin/students/{sid}/class", json={"batch": "BCE09"}, headers=h).status_code == 400  # no reason
    assert client.post(f"/api/admin/students/{sid}/class", json={"batch": "NOPE1", "reason": "x"}, headers=h).status_code == 400
    assert client.post(f"/api/admin/students/{sid}/class", json={"batch": "BCE08", "reason": "x"}, headers=h).status_code == 400
    for user in (world["a"], world["s2"]):
        assert client.post(f"/api/admin/students/{sid}/class", json={"batch": "BCE09", "reason": "x"},
                           headers=auth_header(user)).status_code in (401, 403)
    # The older edit form also logs a class change.
    assert client.put(f"/api/admin/students/{sid}", json={"batch": "BCE09"}, headers=h).status_code == 200
    assert StudentClassChange.query.filter_by(student_id=sid).count() == 1


def test_detailed_report_places_moved_student_in_old_class_for_past_exemptions(client, world):
    h = auth_header(world["admin"])
    # Absence over the last 3 days for BCE09 only; s1 moves into BCE09 today.
    client.post("/api/admin/lecturer-absences", json=_absence_body(world, batches=["BCE09"]), headers=h)
    client.post(f"/api/admin/students/{world['s1'].id}/class", json={"batch": "BCE09", "reason": "Wrong class"}, headers=h)
    AbsenceIndex.invalidate()
    q = f"period=range&from={(TODAY - timedelta(days=3)).isoformat()}&to={(TODAY - timedelta(days=1)).isoformat()}"
    rows = client.get(f"/api/admin/attendance-reports/students?{q}&view=detailed&batch=BCE09", headers=h).json["rows"]
    assert {r["student_id"] for r in rows} == {world["s9"].id}  # s1 was still in BCE08 on those days
