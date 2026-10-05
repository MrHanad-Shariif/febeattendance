"""Course completion report.

Each semester the admin opens a submission period (a start and a deadline).
While it is open, every lecturer reports, for each course they teach to each
class, how many chapters they covered out of those planned; the completion
percentage is worked out from the two. After the deadline submissions close by
themselves, unless the admin reopens the period for chosen lecturers.

Staff holding course_completion:view see every report, filter by semester and
class, and download Excel (the whole semester or one class);
course_completion:edit opens, changes and reopens periods.
"""
import io
from datetime import date

from flask import jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.assignments.routes import lecturer_classes
from app.course_completion import course_completion_bp
from app.extensions import db
from app.models import (
    CompletionPeriod, CompletionReopen, Course, CourseCompletionReport, Timetable, User, utcnow,
)
from app.utils.authz import current_user, permission_required, roles_required
from app.utils.notify import notify, subject_for
from app.utils.permissions import require
from app.utils.rbac import has_permission
from app.utils.settings import get_setting_str
from app.utils.validation import ValidationError, clean_int, clean_text, parse_deadline

EXCEL_HEADERS = ["Semester", "Class Name", "Lecturer Name", "Course Name", "Covered Chapters",
                 "Planned Chapters", "Completion %"]
CLOSED_MESSAGE = ("The course completion report is closed. The deadline has passed, so reports can't be added "
                  "or changed unless the admin reopens it for you.")


def _fmt(dt) -> str:
    return f"{dt:%d %b %Y %H:%M}"


def percent(covered: int, planned: int) -> float:
    return round(covered * 100 / planned, 1) if planned else 0.0


def _current_semester() -> str:
    return get_setting_str("semester_name")


# ---------- Lecturer ----------

def _period_for(user: User) -> CompletionPeriod | None:
    """The period a lecturer works in: one open for them now (newest first),
    else the newest one that has started, else the next one to open."""
    now = utcnow()
    periods = CompletionPeriod.query.order_by(CompletionPeriod.opens_at.desc()).all()
    return (next((p for p in periods if p.is_open_for(user.id, now)), None)
            or next((p for p in periods if p.opens_at <= now), None)
            or (periods[-1] if periods else None))


def _lecturer_view(user: User, period: CompletionPeriod | None) -> dict:
    classes = lecturer_classes(user)
    if period is None:
        return {"period": None, "can_submit": False, "deadline": None, "lecturer_name": user.name,
                "rows": [], "message": "The admin hasn't opened the course completion report yet."}
    now = utcnow()
    reports = {(r.class_name, r.course_name): r for r in period.reports if r.lecturer_id == user.id}
    rows = []
    for c in classes:
        r = reports.pop((c["batch"], c["course_name"]), None)
        rows.append(r.to_dict() if r else {
            "id": None, "semester": period.semester, "class_name": c["batch"], "lecturer_name": user.name,
            "course_name": c["course_name"], "covered_chapters": None, "planned_chapters": None,
            "completion_percent": None,
        })
    # Reports for a class the lecturer no longer teaches stay visible.
    rows += [r.to_dict() for r in reports.values()]
    can_submit = period.is_open_for(user.id, now)
    if can_submit:
        message = None
    elif now < period.opens_at:
        message = f"The course completion report opens on {_fmt(period.opens_at)}."
    else:
        message = CLOSED_MESSAGE
    return {
        "period": {"id": period.id, "semester": period.semester, "opens_at": period.opens_at.isoformat(),
                   "closes_at": period.closes_at.isoformat(), "status": period.status(now)},
        "can_submit": can_submit,
        "deadline": period.deadline_for(user.id).isoformat(),
        "reopened": bool(period.reopened_until(user.id)) and now > period.closes_at,
        "lecturer_name": user.name,
        "rows": sorted(rows, key=lambda r: (r["class_name"], r["course_name"].lower())),
        "message": message,
    }


@course_completion_bp.get("/course-completion/mine")
@roles_required("lecturer")
def my_reports():
    user = current_user()
    return jsonify(_lecturer_view(user, _period_for(user)))


@course_completion_bp.put("/course-completion/mine")
@roles_required("lecturer")
def save_report():
    """Add or change the report for one of the lecturer's courses."""
    user = current_user()
    data = request.get_json(silent=True) or {}
    period = CompletionPeriod.query.get(clean_int(data.get("period_id"), "Period", minimum=1) or 0)
    if period is None:
        raise ValidationError("The course completion report isn't open.", status=404)
    if not period.is_open_for(user.id):
        raise ValidationError(CLOSED_MESSAGE, status=403)

    class_name = clean_text(data.get("class_name"), "Class name", max_len=50, required=True)
    course_name = clean_text(data.get("course_name"), "Course name", max_len=255, required=True)
    report = CourseCompletionReport.query.filter_by(
        period_id=period.id, lecturer_id=user.id, class_name=class_name, course_name=course_name,
    ).first()
    if report is None and {"course_name": course_name, "batch": class_name} not in lecturer_classes(user):
        raise ValidationError("You can only report on courses assigned to you.", status=403)

    planned = clean_int(data.get("planned_chapters"), "Planned chapters", minimum=1, maximum=500)
    covered = clean_int(data.get("covered_chapters"), "Covered chapters", minimum=0, maximum=500)
    if planned is None or covered is None:
        raise ValidationError("Enter both the covered and the planned chapters")
    if covered > planned:
        raise ValidationError("Covered chapters can't be more than the planned chapters")

    now = utcnow()
    if report is None:
        report = CourseCompletionReport(
            period_id=period.id, semester=period.semester, class_name=class_name, lecturer_id=user.id,
            course_name=course_name, submitted_at=now,
        )
        db.session.add(report)
    report.lecturer_name = user.name
    report.covered_chapters = covered
    report.planned_chapters = planned
    report.completion_percent = percent(covered, planned)
    report.updated_at = now
    db.session.commit()
    return jsonify(_lecturer_view(user, period))


# ---------- Admin: periods ----------

def _active_lecturers() -> list[User]:
    return User.query.filter_by(role="lecturer", status="active").order_by(User.name).all()


@course_completion_bp.get("/course-completion/periods")
@permission_required("course_completion:view")
def list_periods():
    periods = CompletionPeriod.query.order_by(CompletionPeriod.opens_at.desc()).all()
    return jsonify({
        "periods": [p.to_dict() for p in periods],
        "current_semester": _current_semester(),
        "lecturers": [{"id": u.id, "name": u.name} for u in _active_lecturers()],
        "can_edit": has_permission(current_user(), "course_completion:edit"),
    })


def _period_dates(data) -> tuple:
    opens_at = parse_deadline(data.get("opens_at"), "Start date")
    closes_at = parse_deadline(data.get("closes_at"), "Deadline")
    if opens_at is None or closes_at is None:
        raise ValidationError("Set both the start date and the deadline")
    if len(str(data.get("opens_at")).strip()) == 10:
        opens_at = opens_at.replace(hour=0, minute=0, second=0)  # a bare date starts at midnight
    if closes_at <= opens_at:
        raise ValidationError("The deadline must be after the start date")
    return opens_at, closes_at


def _announce(period: CompletionPeriod, users: list[User], deadline, reopened: bool = False):
    what = "reopened for you" if reopened else "open"
    notify(
        users, type="course_completion", title=f"Course completion report {what} – {period.semester}",
        message=f"Submit covered and planned chapters for each of your courses by {_fmt(deadline)}.",
        link="/course-completion", related_id=period.id,
        email_subject=subject_for("course_completion_reopened" if reopened else "course_completion_open", period.semester),
        email_lines=[
            f"The course completion report for {period.semester} is {what}.",
            "For each course you teach, enter the chapters covered and the chapters planned.",
            f"Deadline: {_fmt(deadline)}. After it, submissions close automatically.",
        ],
    )


@course_completion_bp.post("/course-completion/periods")
@permission_required("course_completion:edit")
def create_period():
    """Open the report for a semester with a start date and a deadline."""
    data = request.get_json(silent=True) or {}
    semester = clean_text(data.get("semester"), "Semester", max_len=100) or _current_semester()
    if CompletionPeriod.query.filter_by(semester=semester).first():
        raise ValidationError(f"{semester} already has a course completion period. Change its dates instead.", status=409)
    opens_at, closes_at = _period_dates(data)
    if closes_at <= utcnow():
        raise ValidationError("The deadline must be in the future")
    period = CompletionPeriod(semester=semester, opens_at=opens_at, closes_at=closes_at,
                              created_by_id=current_user().id)
    db.session.add(period)
    db.session.flush()
    _announce(period, _active_lecturers(), closes_at)
    db.session.commit()
    return jsonify(period.to_dict()), 201


@course_completion_bp.put("/course-completion/periods/<int:period_id>")
@permission_required("course_completion:edit")
def update_period(period_id):
    """Change the start date or deadline (extending a deadline reopens the
    period for everyone)."""
    period = CompletionPeriod.query.get_or_404(period_id)
    opens_at, closes_at = _period_dates(request.get_json(silent=True) or {})
    period.opens_at, period.closes_at = opens_at, closes_at
    db.session.commit()
    return jsonify(period.to_dict())


@course_completion_bp.post("/course-completion/periods/<int:period_id>/close")
@permission_required("course_completion:edit")
def close_period(period_id):
    """Close submissions now, including any reopenings."""
    period = CompletionPeriod.query.get_or_404(period_id)
    now = utcnow().replace(microsecond=0)
    if period.opens_at > now:
        period.opens_at = now
    period.closes_at = min(period.closes_at, now)
    for r in period.reopenings:
        r.until = min(r.until, now)
    db.session.commit()
    return jsonify(period.to_dict())


@course_completion_bp.post("/course-completion/periods/<int:period_id>/reopen")
@permission_required("course_completion:edit")
def reopen(period_id):
    """Let chosen lecturers add or change their reports until `until`."""
    period = CompletionPeriod.query.get_or_404(period_id)
    data = request.get_json(silent=True) or {}
    ids = data.get("lecturer_ids")
    if not isinstance(ids, list) or not ids or not all(isinstance(i, int) and not isinstance(i, bool) for i in ids):
        raise ValidationError("Choose at least one lecturer")
    until = parse_deadline(data.get("until"), "Reopen until")
    if until is None:
        raise ValidationError("Set the date the reopening ends")
    if until <= utcnow():
        raise ValidationError("The reopening must end in the future")
    lecturers = User.query.filter(User.id.in_(set(ids)), User.role == "lecturer").all()
    if len(lecturers) != len(set(ids)):
        raise ValidationError("Choose lecturers from the list")

    user = current_user()
    for lecturer in lecturers:
        row = next((r for r in period.reopenings if r.lecturer_id == lecturer.id), None)
        if row is None:
            row = CompletionReopen(lecturer_id=lecturer.id)
            period.reopenings.append(row)
        row.until, row.granted_by_id, row.granted_at = until, user.id, utcnow()
    _announce(period, lecturers, until, reopened=True)
    db.session.commit()
    return jsonify(period.to_dict())


@course_completion_bp.delete("/course-completion/periods/<int:period_id>/reopen/<int:lecturer_id>")
@permission_required("course_completion:edit")
def revoke_reopen(period_id, lecturer_id):
    period = CompletionPeriod.query.get_or_404(period_id)
    row = CompletionReopen.query.filter_by(period_id=period.id, lecturer_id=lecturer_id).first_or_404()
    db.session.delete(row)
    db.session.commit()
    db.session.refresh(period)
    return jsonify(period.to_dict())


# ---------- Admin: reports ----------

def _expected_courses(class_name: str | None) -> list[tuple[str, str, User | None]]:
    """(class, course, lecturer) taught this semester, from the course
    catalogue and the timetable."""
    courses = Course.query.filter(Course.batch.isnot(None), Course.lecturer_id.isnot(None))
    sessions = Timetable.query.filter(Timetable.batch.isnot(None))
    if class_name:
        courses = courses.filter(Course.batch == class_name)
        sessions = sessions.filter(Timetable.batch == class_name)
    seen: dict[tuple[str, str, int], User] = {}
    for c in courses.all():
        seen.setdefault((c.batch, c.name, c.lecturer_id), c.lecturer)
    for t in sessions.all():
        seen.setdefault((t.batch, t.course_name, t.lecturer_id), t.lecturer)
    return [(batch, course, lecturer) for (batch, course, _), lecturer in seen.items()]


def _report_rows(semester: str, class_name: str | None) -> list[dict]:
    query = CourseCompletionReport.query.filter_by(semester=semester)
    if class_name:
        query = query.filter_by(class_name=class_name)
    rows = [dict(r.to_dict(), submitted=True) for r in query.all()]
    if semester == _current_semester():
        # Show what's still missing for the current semester.
        done = {(r["class_name"], r["course_name"], r["lecturer_id"]) for r in rows}
        for batch, course, lecturer in _expected_courses(class_name):
            lecturer_id = lecturer.id if lecturer else None
            if (batch, course, lecturer_id) not in done:
                rows.append({
                    "id": None, "semester": semester, "class_name": batch, "lecturer_id": lecturer_id,
                    "lecturer_name": lecturer.name if lecturer else "—", "course_name": course,
                    "covered_chapters": None, "planned_chapters": None, "completion_percent": None,
                    "submitted": False,
                })
    return sorted(rows, key=lambda r: (r["class_name"], r["course_name"].lower(), r["lecturer_name"]))


def _filters() -> tuple[str, str | None]:
    semester = clean_text(request.args.get("semester"), "Semester", max_len=100) or _current_semester()
    class_name = clean_text(request.args.get("class_name"), "Class", max_len=50)
    return semester, class_name


@course_completion_bp.get("/course-completion/reports")
@permission_required("course_completion:view")
def list_reports():
    semester, class_name = _filters()
    semesters = {p.semester for p in CompletionPeriod.query.all()} | {_current_semester()}
    all_rows = _report_rows(semester, None)
    rows = [r for r in all_rows if not class_name or r["class_name"] == class_name]
    return jsonify({
        "semester": semester,
        "class_name": class_name,
        "semesters": sorted(semesters, reverse=True),
        "classes": sorted({r["class_name"] for r in all_rows}),
        "rows": rows,
    })


def _excel_text(value):
    """Cells are data, never formulas: a typed '=...' would otherwise run in Excel."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _safe_filename(text: str) -> str:
    return "".join(ch if ch.isalnum() else "-" for ch in text).strip("-")[:60] or "report"


@course_completion_bp.get("/course-completion/export.xlsx")
@permission_required("course_completion:view")
def export_excel():
    """Download All (no class) or Download by Class (class_name=...)."""
    semester, class_name = _filters()
    rows = [r for r in _report_rows(semester, class_name) if r["submitted"]]

    wb = Workbook()
    ws = wb.active
    ws.title = (class_name or "All classes")[:31]
    ws.append(EXCEL_HEADERS)
    for i, w in enumerate([28, 14, 28, 38, 18, 18, 15], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="15803D")
    for r in rows:
        ws.append([_excel_text(v) for v in (
            r["semester"], r["class_name"], r["lecturer_name"], r["course_name"],
            r["covered_chapters"], r["planned_chapters"], r["completion_percent"] / 100,
        )])
        ws.cell(row=ws.max_row, column=7).number_format = "0.0%"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"course-completion-{_safe_filename(semester)}-{_safe_filename(class_name) if class_name else 'all'}-{date.today():%Y-%m-%d}.xlsx"
    return send_file(buf, as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
