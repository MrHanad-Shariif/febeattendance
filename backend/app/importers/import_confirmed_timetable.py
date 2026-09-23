"""Import the confirmed (green-highlighted) class sessions from the faculty's
detailed per-batch timetable workbook ("Time_Table_OCTOBER 2026 - FEB 2027.xlsx").

Each sheet in that workbook is one student batch (e.g. "BCE08"), laid out as a
Day / Time / Instructor / Subject / ROOM grid. Only rows where the Instructor
or Subject cell is filled with one of the faculty's "confirmed" green shades
are imported; rows with no instructor assigned, or with a subject but no
green highlight, are left out on purpose (they are not finalized yet).

Back-to-back time slots for the same lecturer/course/room/day are merged into
a single timetable row with a `sessions` count, and the same lecturer/course/
room/time block recurring on multiple days is combined into one row with a
comma-separated day list -- matching how the Timetable table represents a
class.

This is a one-time, authoritative import: it REPLACES every existing
Timetable row (which, before this, only had bare lecturer/course pairs with
no schedule) with the confirmed rows parsed here.

Usage (from backend/, with the venv active):
    flask import-confirmed-timetable "..\\Time_Table_OCTOBER 2026 - FEB 2027 (1).xlsx"
"""
import re
import difflib
import datetime as dt

import openpyxl

from app.extensions import db
from app.models import User, Timetable

SKIP_SHEETS = {"ROOM MNGT"}

GREEN_FILLS = {"FF92D050", "FFC2D69B", "FF6AA84F", "FFD6E3BC"}

DAY_NORMALIZE = {
    "sat": "Sat", "saturday": "Sat",
    "sun": "Sun", "sunday": "Sun",
    "mon": "Mon", "monday": "Mon",
    "tue": "Tue", "tues": "Tue", "tuesday": "Tue",
    "wed": "Wed", "wednesday": "Wed",
    "thu": "Thu", "thur": "Thu", "thurs": "Thu", "thursday": "Thu", "thus": "Thu",
    "fri": "Fri", "friday": "Fri",
}

TIME_RE = re.compile(r"(\d{1,2}):(\d{2})\s*([AaPp][Mm])?")

# Typo/shorthand corrections confirmed against the existing lecturer roster
# (see the session that built this importer for how each was verified).
NAME_CORRECTIONS = {
    "mahdi yussuf abdi": "Mahdi Yusuf Abdi",
    "maryam abdirahman mahmud": "Maryam Abdirahman Mahamud",
    "zakaria ibrahim ali": "Zakarie ibrahim Ali",
    "amina ahmed hasssan nur": "Amina Ahmed Hassan Nur",
    "liban ali sahal": "Liban Sahal",
    "dr.mohamed hussein": "Dr-Mohamed Hussein Ibrahim",
    "dr.mohamud omar": "Dr. Mohamud Omar Aden",
    "anas moallim": "Anas Hussein Mohamud",
    "dr. mohamed abukar": "Dr-Mohamed Abubakar Hagi-Mohamed",
}


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.strip().lower())


def parse_time_token(hour, minute, ampm):
    hour = int(hour)
    minute = int(minute)
    if 1 <= hour <= 7 and ampm and ampm.lower() == "pm":
        hour += 12
    return dt.time(hour % 24, minute)


def parse_time_range(value):
    if isinstance(value, dt.time):
        return value, None
    if isinstance(value, dt.datetime):
        return value.time(), None
    if not isinstance(value, str):
        return None, None
    matches = TIME_RE.findall(value)
    if len(matches) < 2:
        return None, None
    return parse_time_token(*matches[0]), parse_time_token(*matches[1])


def is_green(cell):
    fill = cell.fill
    if fill and fill.patternType == "solid":
        rgb = fill.fgColor.rgb
        if isinstance(rgb, str) and rgb in GREEN_FILLS:
            return True
    return False


def find_header(ws):
    for row in ws.iter_rows(min_row=1, max_row=min(15, ws.max_row)):
        texts = {}
        for c in row:
            if c.value is not None:
                texts[str(c.value).strip().lower()] = c.column
        if "day" in texts:
            return row[0].row, texts
    return None, None


def format_room(value):
    if value is None or value == "":
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_sheet(ws, batch_name):
    header_row, col_map = find_header(ws)
    if header_row is None:
        return []
    day_col = col_map.get("day")
    time_col = col_map.get("time")
    instr_col = col_map.get("instructor")
    subj_col = col_map.get("subject")
    room_col = col_map.get("room")
    if not all([day_col, time_col, instr_col, subj_col]):
        return []

    records = []
    current_day = None
    for row_idx in range(header_row + 1, ws.max_row + 1):
        day_cell = ws.cell(row=row_idx, column=day_col)
        if day_cell.value:
            raw = str(day_cell.value).strip().lower()
            current_day = DAY_NORMALIZE.get(raw, str(day_cell.value).strip())

        time_cell = ws.cell(row=row_idx, column=time_col)
        instr_cell = ws.cell(row=row_idx, column=instr_col)
        subj_cell = ws.cell(row=row_idx, column=subj_col)
        room_cell = ws.cell(row=row_idx, column=room_col) if room_col else None

        if not time_cell.value or not instr_cell.value:
            continue
        if not (is_green(instr_cell) or is_green(subj_cell)):
            continue

        start, end = parse_time_range(time_cell.value)
        if not start or not end:
            continue

        lecturer_name = str(instr_cell.value).strip()
        course_name = str(subj_cell.value).strip() if subj_cell.value else "Unknown"
        room = format_room(room_cell.value) if room_cell else None

        records.append({
            "batch": batch_name, "day": current_day, "start": start, "end": end,
            "lecturer": lecturer_name, "course": course_name, "room": room,
        })
    return records


def merge_contiguous(records):
    merged = []
    for rec in records:
        if merged:
            last = merged[-1]
            if (last["batch"] == rec["batch"] and last["lecturer"] == rec["lecturer"]
                    and last["course"] == rec["course"] and last["room"] == rec["room"]
                    and last["day"] == rec["day"] and last["end"] == rec["start"]):
                last["end"] = rec["end"]
                last["sessions"] += 1
                continue
        rec = dict(rec)
        rec["sessions"] = 1
        merged.append(rec)
    return merged


def combine_days(merged):
    groups, order = {}, []
    for rec in merged:
        key = (rec["batch"], rec["lecturer"], rec["course"], rec["room"], rec["start"], rec["end"], rec["sessions"])
        if key not in groups:
            groups[key] = {**rec, "days": []}
            order.append(key)
        if rec["day"] and rec["day"] not in groups[key]["days"]:
            groups[key]["days"].append(rec["day"])
    return [groups[k] for k in order]


def is_subset_match(short, full):
    short_tokens, full_tokens = norm(short).split(), norm(full).split()
    it = iter(full_tokens)
    return all(any(tok == f for f in it) for tok in short_tokens)


def resolve_lecturer_name(raw_name, existing_names):
    n = norm(raw_name)
    for existing in existing_names:
        if norm(existing) == n:
            return existing
    if n in NAME_CORRECTIONS:
        corrected = NAME_CORRECTIONS[n]
        for existing in existing_names:
            if norm(existing) == norm(corrected):
                return existing
    subset_matches = [e for e in existing_names if is_subset_match(raw_name, e)]
    if len(subset_matches) == 1:
        return subset_matches[0]
    return None  # no confident match -> caller creates a new lecturer


def import_confirmed_timetable(xlsx_path: str):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    existing_lecturers = User.query.filter_by(role="lecturer").all()
    existing_names = [u.name for u in existing_lecturers]
    known_names_normalized = {norm(n) for n in existing_names} | set(NAME_CORRECTIONS.keys())

    all_final = []
    for name in wb.sheetnames:
        if name in SKIP_SHEETS:
            continue
        ws = wb[name]
        raw = parse_sheet(ws, name.strip())

        # Fix rows where the Instructor/Subject columns were swapped in the
        # source sheet: the "lecturer" value isn't a known name, but the
        # "course" value is.
        for rec in raw:
            if norm(rec["lecturer"]) not in known_names_normalized and norm(rec["course"]) in known_names_normalized:
                rec["lecturer"], rec["course"] = rec["course"], rec["lecturer"]

        all_final.extend(combine_days(merge_contiguous(raw)))

    by_name = {u.name: u for u in existing_lecturers}
    created_lecturers = 0
    resolved_rows = []
    unresolved_names = set()

    for rec in all_final:
        canonical = resolve_lecturer_name(rec["lecturer"], existing_names)
        if canonical:
            lecturer = by_name[canonical]
        else:
            lecturer = by_name.get(rec["lecturer"])
            if not lecturer:
                lecturer = User(
                    name=rec["lecturer"].strip(),
                    email=f"pending+{abs(hash(rec['lecturer']))}@example.invalid",
                    role="lecturer", status="invited",
                )
                db.session.add(lecturer)
                db.session.flush()
                by_name[lecturer.name] = lecturer
                existing_names.append(lecturer.name)
                created_lecturers += 1
                unresolved_names.add(rec["lecturer"])
        resolved_rows.append((lecturer, rec))

    # This is an authoritative re-import: clear the old placeholder timetable
    # (bare lecturer/course pairs with no schedule) before inserting the real one.
    deleted = Timetable.query.delete()

    for lecturer, rec in resolved_rows:
        db.session.add(Timetable(
            lecturer_id=lecturer.id,
            course_name=rec["course"],
            batch=rec["batch"],
            room=rec["room"],
            days=",".join(rec["days"]) if rec["days"] else None,
            start_time=rec["start"],
            end_time=rec["end"],
            sessions=rec["sessions"],
        ))

    db.session.commit()

    print(f"Deleted {deleted} old placeholder timetable rows.")
    print(f"Inserted {len(resolved_rows)} confirmed timetable rows across {len(wb.sheetnames) - len(SKIP_SHEETS)} batches.")
    print(f"Created {created_lecturers} new lecturer(s) not found in the existing roster: {sorted(unresolved_names)}")
