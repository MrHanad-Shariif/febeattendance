"""Special exam registration.

A student who missed an exam registers for a special exam: name, batch,
course, reason, the date of the missed exam, the shift (1 or 2), phone and
ID number, and may attach supporting documents (a medical certificate, a
letter...). Registration is closed unless the Administration Team has set a
registration period (start and end) and the current time is inside it.

The Administration Team (or a staff role holding special_exams:view /
special_exams:edit) sees every request, approves or declines it with an
optional note, and downloads the list as Excel. The student sees the request
as pending until then, and is emailed the decision.
"""
import io
import re
from datetime import date

from flask import jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.extensions import db, limiter
from app.models import (
    Course, SpecialExamDocument, SpecialExamPeriod, SpecialExamRequest, Timetable, User, utcnow,
)
from app.special_exams import special_exams_bp
from app.utils.authz import current_user, roles_required
from app.utils.notify import notify, subject_for
from app.utils.permissions import can_decide_special_exams, can_view_special_exams, require
from app.utils.uploads import delete_document, save_document, send_document
from app.utils.validation import ValidationError, clean_text, parse_deadline, parse_iso_date

PHONE_RE = re.compile(r"^\+?[0-9][0-9 \-]{5,18}[0-9]$")
STATUS_LABELS = {"pending": "Pending", "approved": "Approved", "declined": "Declined"}
MAX_DOCUMENTS = 5
CLOSED_MESSAGE = "Special exam registration is closed. You can register only during the period set by the Administration Team."


def _payload():
    if request.files or request.form:
        return request.form
    return request.get_json(silent=True) or {}


# ---------- Registration period ----------

def open_period() -> SpecialExamPeriod | None:
    now = utcnow()
    return (SpecialExamPeriod.query
            .filter(SpecialExamPeriod.opens_at <= now, SpecialExamPeriod.closes_at >= now)
            .order_by(SpecialExamPeriod.closes_at.desc()).first())


def current_period() -> SpecialExamPeriod | None:
    """The open period, else the next one to open, else the last one that ran."""
    now = utcnow()
    return (open_period()
            or SpecialExamPeriod.query.filter(SpecialExamPeriod.opens_at > now).order_by(SpecialExamPeriod.opens_at).first()
            or SpecialExamPeriod.query.order_by(SpecialExamPeriod.closes_at.desc()).first())


def _period_or_none(period: SpecialExamPeriod | None):
    return period.to_dict() if period else None


def _batch_courses() -> dict[str, list[str]]:
    """{batch: [course names]} from the course catalogue and the timetable."""
    pairs = set(
        Course.query.with_entities(Course.batch, Course.name).filter(Course.batch.isnot(None)).distinct()
    ) | set(
        Timetable.query.with_entities(Timetable.batch, Timetable.course_name).filter(Timetable.batch.isnot(None)).distinct()
    )
    result: dict[str, set[str]] = {}
    for batch, course in pairs:
        if batch and course:
            result.setdefault(batch, set()).add(course)
    return {b: sorted(result[b], key=str.lower) for b in sorted(result)}


# ---------- Student ----------

@special_exams_bp.get("/special-exams/options")
@roles_required("student")
def options():
    user = current_user()
    return jsonify({
        "batches": _batch_courses(),
        "defaults": {"full_name": user.name, "batch": user.batch, "id_number": user.student_id_number},
        "period": _period_or_none(current_period()),
    })


@special_exams_bp.get("/special-exams/mine")
@roles_required("student")
def my_requests():
    rows = (SpecialExamRequest.query.filter_by(student_id=current_user().id)
            .order_by(SpecialExamRequest.created_at.desc()).all())
    return jsonify([r.to_dict() for r in rows])


@special_exams_bp.post("/special-exams")
@roles_required("student")
@limiter.limit("20/hour", key_func=lambda: f"user:{current_user().id}")
def create_request():
    user = current_user()
    period = open_period()
    if period is None:
        raise ValidationError(CLOSED_MESSAGE, status=403)
    data = _payload()
    uploads = [f for f in request.files.getlist("documents") if f and f.filename]
    if len(uploads) > MAX_DOCUMENTS:
        raise ValidationError(f"You can attach at most {MAX_DOCUMENTS} documents")

    full_name = clean_text(data.get("full_name"), "Name", max_len=200, required=True)
    batch = clean_text(data.get("batch"), "Batch", max_len=50, required=True)
    course_name = clean_text(data.get("course_name"), "Course", max_len=255, required=True)
    reason = clean_text(data.get("reason"), "Reason", max_len=2000, required=True)
    phone = clean_text(data.get("phone"), "Phone number", max_len=30, required=True)
    id_number = clean_text(data.get("id_number"), "ID number", max_len=50, required=True)
    if data.get("exam_date") in (None, ""):
        raise ValidationError("Exam date is required")
    exam_date = parse_iso_date(data.get("exam_date"), "Exam date")

    courses = _batch_courses()
    if batch not in courses:
        raise ValidationError("Choose a batch from the list")
    if course_name not in courses[batch]:
        raise ValidationError("Choose a course from the list for that batch")
    if exam_date > date.today():
        raise ValidationError("The exam date can't be in the future: enter the date the exam was held")
    if str(data.get("shift")) not in ("1", "2"):
        raise ValidationError("Shift must be 1 or 2")
    if not PHONE_RE.match(phone):
        raise ValidationError("Enter a valid phone number (digits, spaces or dashes, optionally starting with +)")

    duplicate = SpecialExamRequest.query.filter_by(
        student_id=user.id, course_name=course_name, exam_date=exam_date, status="pending",
    ).first()
    if duplicate:
        return jsonify({"error": "You already have a pending request for this course and exam date."}), 409

    row = SpecialExamRequest(
        student_id=user.id, period_id=period.id, full_name=full_name, batch=batch, course_name=course_name,
        reason=reason, exam_date=exam_date, shift=int(data["shift"]), phone=phone, id_number=id_number,
    )
    saved_paths = []
    try:
        for upload in uploads:
            name, path = save_document(upload, "special_exams")
            saved_paths.append(path)
            row.documents.append(SpecialExamDocument(file_name=name, file_path=path, size_bytes=_size_of(path)))
    except ValidationError as err:
        for path in saved_paths:
            delete_document(path)
        raise ValidationError(f"{upload.filename}: {err.message}")
    db.session.add(row)
    db.session.flush()

    # In-app heads-up for reviewers only; no email, to keep inboxes quiet.
    reviewers = [u for u in User.query.filter(User.role != "student", User.status == "active")
                 if can_decide_special_exams(u)]
    notify(reviewers, type="special_exam_request", title=f"Special exam request – {full_name}",
           message=f"{course_name} ({batch}), exam of {exam_date:%d %b %Y}, shift {row.shift}.",
           link="/admin/special-exams", related_id=row.id)
    db.session.commit()
    return jsonify(row.to_dict()), 201


def _size_of(relative_path: str) -> int | None:
    import os

    from flask import current_app
    try:
        return os.path.getsize(os.path.join(current_app.config["UPLOAD_FOLDER"], relative_path))
    except OSError:
        return None


@special_exams_bp.get("/special-exams/<int:request_id>/documents/<int:document_id>")
@roles_required("admin", "lecturer", "student")
def download_document(request_id, document_id):
    """The student who sent the request, or a reviewer."""
    user = current_user()
    row = SpecialExamRequest.query.get_or_404(request_id)
    require(row.student_id == user.id if user.role == "student" else can_view_special_exams(user))
    doc = SpecialExamDocument.query.filter_by(id=document_id, request_id=row.id).first_or_404()
    return send_document(doc.file_path, doc.file_name, inline=request.args.get("inline") == "1")


# ---------- Administration Team ----------

@special_exams_bp.get("/special-exams/period")
@roles_required("admin", "lecturer")
def get_period():
    require(can_view_special_exams(current_user()))
    history = SpecialExamPeriod.query.order_by(SpecialExamPeriod.opens_at.desc()).limit(20).all()
    return jsonify({"current": _period_or_none(current_period()), "history": [p.to_dict() for p in history]})


def _period_fields(data) -> tuple[str | None, object, object]:
    title = clean_text(data.get("title"), "Title", max_len=200)
    opens_at = parse_deadline(data.get("opens_at"), "Start")
    closes_at = parse_deadline(data.get("closes_at"), "End")
    if opens_at is None or closes_at is None:
        raise ValidationError("Set both the start and the end of the registration period")
    if len(str(data.get("opens_at")).strip()) == 10:
        opens_at = opens_at.replace(hour=0, minute=0, second=0)  # a bare date starts at midnight
    if closes_at <= opens_at:
        raise ValidationError("The end must be after the start")
    return title, opens_at, closes_at


@special_exams_bp.post("/special-exams/period")
@roles_required("admin", "lecturer")
def create_period():
    """Open a new registration period (start and end set by the Administration Team)."""
    user = current_user()
    require(can_decide_special_exams(user), "Only the Administration Team can set the registration period.")
    title, opens_at, closes_at = _period_fields(request.get_json(silent=True) or {})
    if closes_at <= utcnow():
        raise ValidationError("The end of the registration period must be in the future")
    overlap = SpecialExamPeriod.query.filter(
        SpecialExamPeriod.opens_at <= closes_at, SpecialExamPeriod.closes_at >= opens_at,
    ).first()
    if overlap:
        raise ValidationError("This overlaps another registration period. Edit or close that one instead.", status=409)
    period = SpecialExamPeriod(title=title, opens_at=opens_at, closes_at=closes_at, created_by_id=user.id)
    db.session.add(period)
    db.session.commit()
    return jsonify(period.to_dict()), 201


@special_exams_bp.put("/special-exams/period/<int:period_id>")
@roles_required("admin", "lecturer")
def update_period(period_id):
    require(can_decide_special_exams(current_user()), "Only the Administration Team can set the registration period.")
    period = SpecialExamPeriod.query.get_or_404(period_id)
    title, opens_at, closes_at = _period_fields(request.get_json(silent=True) or {})
    overlap = SpecialExamPeriod.query.filter(
        SpecialExamPeriod.id != period.id,
        SpecialExamPeriod.opens_at <= closes_at, SpecialExamPeriod.closes_at >= opens_at,
    ).first()
    if overlap:
        raise ValidationError("This overlaps another registration period.", status=409)
    period.title, period.opens_at, period.closes_at = title, opens_at, closes_at
    db.session.commit()
    return jsonify(period.to_dict())


@special_exams_bp.post("/special-exams/period/<int:period_id>/close")
@roles_required("admin", "lecturer")
def close_period(period_id):
    """End registration now (an upcoming period is cancelled the same way)."""
    require(can_decide_special_exams(current_user()), "Only the Administration Team can set the registration period.")
    period = SpecialExamPeriod.query.get_or_404(period_id)
    now = utcnow().replace(microsecond=0)
    if period.status(now) == "closed":
        raise ValidationError("This registration period has already ended")
    if period.opens_at > now:
        period.opens_at = now
    period.closes_at = now
    db.session.commit()
    return jsonify(period.to_dict())

def _filtered_query():
    query = SpecialExamRequest.query
    status = request.args.get("status")
    if status:
        if status not in SpecialExamRequest.STATUSES:
            raise ValidationError("Status must be pending, approved or declined")
        query = query.filter_by(status=status)
    return query.order_by(SpecialExamRequest.created_at.desc())


@special_exams_bp.get("/special-exams")
@roles_required("admin", "lecturer")
def list_requests():
    require(can_view_special_exams(current_user()))
    return jsonify([r.to_dict() for r in _filtered_query().all()])


@special_exams_bp.post("/special-exams/<int:request_id>/decision")
@roles_required("admin", "lecturer")
def decide(request_id):
    user = current_user()
    require(can_decide_special_exams(user), "Only the Administration Team can approve or decline requests.")
    row = SpecialExamRequest.query.get_or_404(request_id)
    data = request.get_json(silent=True) or {}

    decision = data.get("decision")
    if decision not in ("approved", "declined"):
        raise ValidationError("Decision must be approved or declined")
    note = clean_text(data.get("note"), "Reason for the decision", max_len=1000)

    row.status = decision
    row.decision_note = note
    row.decided_by_id = user.id
    row.decided_at = utcnow()

    word = "approved" if decision == "approved" else "declined"
    lines = [
        f"Your special exam request has been {word}.",
        f"Course: {row.course_name} ({row.batch})",
        f"Exam date: {row.exam_date:%d %b %Y}, shift {row.shift}",
    ]
    if note:
        lines.append(f"Note from the Administration Team: {note}")
    notify([row.student], type="special_exam_decision", include_students=True,
           title=f"Your special exam request was {word}", message=lines[0] + (f" {note}" if note else ""),
           link="/special-exams", related_id=row.id,
           email_subject=subject_for(f"special_exam_{decision}", row.course_name), email_lines=lines)
    db.session.commit()
    return jsonify(row.to_dict())


def _excel_text(value):
    """Cells are data, never formulas: a typed '=...' would otherwise run in Excel."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


@special_exams_bp.get("/special-exams/export.xlsx")
@roles_required("admin", "lecturer")
def export_excel():
    require(can_view_special_exams(current_user()))
    rows = _filtered_query().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Special exams"
    headers = ["No.", "Submitted", "Student name", "ID number", "Batch", "Course", "Exam date", "Shift",
               "Phone", "Email", "Reason", "Documents", "Status", "Decision note", "Decided by", "Decided on"]
    widths = [6, 17, 28, 14, 10, 34, 12, 7, 16, 28, 45, 30, 11, 40, 22, 17]
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="15803D")

    fills = {"approved": "DCFCE7", "declined": "FEE2E2", "pending": "FEF9C3"}
    for n, r in enumerate(rows, start=1):
        ws.append([_excel_text(v) for v in (
            n, r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "", r.full_name, r.id_number,
            r.batch, r.course_name, r.exam_date, r.shift, r.phone, r.student.email if r.student else "",
            r.reason, "\n".join(d.file_name for d in r.documents), STATUS_LABELS[r.status], r.decision_note or "",
            r.decided_by.name if r.decided_by else "", r.decided_at.strftime("%Y-%m-%d %H:%M") if r.decided_at else "",
        )])
        ws.cell(row=n + 1, column=7).number_format = "yyyy-mm-dd"
        ws.cell(row=n + 1, column=13).fill = PatternFill("solid", fgColor=fills[r.status])
        for col in (11, 12, 14):
            ws.cell(row=n + 1, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    status = request.args.get("status") or "all"
    return send_file(
        buf, as_attachment=True, download_name=f"special-exam-requests-{status}-{date.today():%Y-%m-%d}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
